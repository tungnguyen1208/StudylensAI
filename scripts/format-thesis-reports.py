from __future__ import annotations

from io import BytesIO
import importlib.util
from pathlib import Path
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SRS_PATH = ROOT / "SRS_StudyLens_AI_Persistent_Activation_v1.2 (2).docx"
SDS_PATH = ROOT / "SDS - StudyLens AI - Học tập có tương tác với Youtube.docx"

FONT_NAME = "Times New Roman"

LOGICAL_DOT = r'''
digraph logical_architecture {
  graph [bgcolor="white", pad="0.25", ranksep="0.48", nodesep="0.30", fontname="Arial"];
  node [shape=box, style="rounded,filled", fontname="Arial", fontsize=11,
        color="#3F6B47", fillcolor="#F2F8EE", margin="0.16,0.10"];
  edge [color="#4F7655", fontcolor="#38563D", fontname="Arial", fontsize=9, penwidth=1.25];

  subgraph cluster_presentation {
    label="Presentation Layer"; color="#A8C995"; style="rounded,dashed";
    panel [label="Side Panel UI"];
    context [label="YouTube Context Adapter"];
  }
  subgraph cluster_application {
    label="Application Layer"; color="#A8C995"; style="rounded,dashed";
    activation [label="Activation Management"];
    learning [label="Session & Learning Package"];
    assessment [label="Assessment & History"];
  }
  subgraph cluster_integration {
    label="Integration Layer"; color="#A8C995"; style="rounded,dashed";
    api [label="ASP.NET Core REST API"];
    ai [label="FastAPI AI Gateway"];
  }
  subgraph cluster_data {
    label="Data & External Services"; color="#A8C995"; style="rounded,dashed";
    db [shape=cylinder, label="PostgreSQL"];
    youtube [shape=ellipse, fillcolor="#FFF8DB", color="#9A7B22", label="YouTube Timedtext"];
    gemini [shape=ellipse, fillcolor="#FFF8DB", color="#9A7B22", label="Google Gemini API"];
  }

  context -> activation [label="video context"];
  panel -> learning [dir=both, label="commands / state"];
  activation -> api;
  learning -> api;
  assessment -> api;
  context -> youtube [label="full captions"];
  api -> db [label="persistence"];
  api -> ai [label="AI requests"];
  ai -> gemini [label="HTTPS"];
}
'''

PHYSICAL_DOT = r'''
digraph physical_architecture {
  graph [bgcolor="white", pad="0.25", ranksep="0.52", nodesep="0.32", fontname="Arial"];
  node [shape=box, style="rounded,filled", fontname="Arial", fontsize=11,
        color="#3F6B47", fillcolor="#F2F8EE", margin="0.16,0.10"];
  edge [color="#4F7655", fontcolor="#38563D", fontname="Arial", fontsize=9, penwidth=1.25];

  subgraph cluster_device {
    label="User Device"; color="#A8C995"; style="rounded,dashed";
    browser [label="Chrome / Edge"];
    extension [label="StudyLens Extension\nManifest V3"];
    browser -> extension [label="Side Panel"];
  }

  subgraph cluster_host {
    label="StudyLens Server (Docker Host)"; color="#A8C995"; style="rounded,dashed";
    api [label="API Container\nASP.NET Core 8\nHost: 127.0.0.1:5000"];
    migrate [label="Migration Job\n--migrate-only"];
    ai [label="AI Container\nPython 3.12 + FastAPI\nInternal :8000"];
    db [shape=cylinder, label="PostgreSQL 16\nInternal :5432"];
    volume [shape=cylinder, fillcolor="#E8F2FF", color="#496B8A", label="Named Data Volume"];
    api -> ai [label="Docker network"];
    api -> db [label="EF Core"];
    migrate -> db [label="schema migration"];
    db -> volume [label="persistent data"];
  }

  youtube [shape=ellipse, fillcolor="#FFF8DB", color="#9A7B22", label="YouTube Web"];
  gemini [shape=ellipse, fillcolor="#FFF8DB", color="#9A7B22", label="Google Gemini API"];

  extension -> api [label="HTTP contract 0.5.0"];
  extension -> youtube [label="page context / captions"];
  ai -> gemini [label="HTTPS"];
}
'''


def set_run_font(run, size: float | None = None, *, bold: bool | None = None,
                 italic: bool | None = None) -> None:
    run.font.name = FONT_NAME
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT_NAME)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, size: float, *, bold: bool = False) -> None:
    style.font.name = FONT_NAME
    style.font.size = Pt(size)
    style.font.bold = bold
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT_NAME)


def ensure_style(document: Document, name: str, *, size: float, bold: bool = False,
                 italic: bool = False, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    styles = document.styles
    style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(style, size, bold=bold)
    style.font.italic = italic
    style.paragraph_format.alignment = alignment
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.line_spacing = 1.2
    return style


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def set_text(paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run)


def add_page_field(paragraph) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    set_run_font(run, 13)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_sections(document: Document) -> None:
    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)
        section.header_distance = Cm(1.2)
        section.footer_distance = Cm(1.2)
        section.different_first_page_header_footer = False
        header = section.header
        if not header.paragraphs:
            header.add_paragraph()
        add_page_field(header.paragraphs[0])
        for extra in header.paragraphs[1:]:
            clear_paragraph(extra)
        for paragraph in section.footer.paragraphs:
            clear_paragraph(paragraph)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    set_style_font(normal, 13)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    heading1 = document.styles["Heading 1"]
    set_style_font(heading1, 18, bold=True)
    heading1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading1.paragraph_format.first_line_indent = Cm(0)
    heading1.paragraph_format.space_before = Pt(0)
    heading1.paragraph_format.space_after = Pt(12)
    heading1.paragraph_format.page_break_before = True
    heading1.paragraph_format.keep_with_next = True

    heading2 = document.styles["Heading 2"]
    set_style_font(heading2, 16, bold=True)
    heading2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading2.paragraph_format.first_line_indent = Cm(0)
    heading2.paragraph_format.space_before = Pt(6)
    heading2.paragraph_format.space_after = Pt(6)
    heading2.paragraph_format.keep_with_next = True

    heading3 = document.styles["Heading 3"]
    set_style_font(heading3, 14, bold=True)
    heading3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading3.paragraph_format.first_line_indent = Cm(0)
    heading3.paragraph_format.space_before = Pt(6)
    heading3.paragraph_format.space_after = Pt(6)
    heading3.paragraph_format.keep_with_next = True

    ensure_style(document, "DATN Table Caption", size=13, bold=True,
                 alignment=WD_ALIGN_PARAGRAPH.CENTER)
    ensure_style(document, "DATN Figure Caption", size=13, italic=True,
                 alignment=WD_ALIGN_PARAGRAPH.CENTER)
    figure_note = ensure_style(document, "DATN Figure Note", size=13,
                               alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    figure_note.paragraph_format.space_before = Pt(3)
    figure_note.paragraph_format.space_after = Pt(6)
    ensure_style(document, "DATN Note", size=13, italic=True,
                 alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)

    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        if name in document.styles:
            set_style_font(document.styles[name], 13)
            document.styles[name].paragraph_format.line_spacing = 1.2


def paragraph_after(paragraph, text: str = "", style: str | None = None):
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    new_paragraph = paragraph.__class__(element, paragraph._parent)
    if style:
        new_paragraph.style = style
    if text:
        run = new_paragraph.add_run(text)
        set_run_font(run)
    return new_paragraph


def paragraph_before_table(table, text: str, style: str):
    element = OxmlElement("w:p")
    table._tbl.addprevious(element)
    paragraph = Paragraph(element, table._parent)
    paragraph.style = style
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def paragraph_after_table(table, text: str, style: str):
    element = OxmlElement("w:p")
    table._tbl.addnext(element)
    paragraph = Paragraph(element, table._parent)
    paragraph.style = style
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def previous_paragraph_text(table) -> str:
    node = table._tbl.getprevious()
    while node is not None:
        if node.tag == qn("w:p"):
            return "".join(node.itertext()).strip()
        node = node.getprevious()
    return ""


def next_paragraph_text(table) -> str:
    node = table._tbl.getnext()
    while node is not None:
        if node.tag == qn("w:p"):
            return "".join(node.itertext()).strip()
        node = node.getnext()
    return ""


def add_table_caption(table, text: str) -> None:
    previous = previous_paragraph_text(table)
    if previous.startswith("Bảng "):
        node = table._tbl.getprevious()
        while node is not None and node.tag != qn("w:p"):
            node = node.getprevious()
        if node is not None:
            paragraph = Paragraph(node, table._parent)
            set_text(paragraph, text)
            paragraph.style = "DATN Table Caption"
        return
    paragraph_before_table(table, text, "DATN Table Caption")


def add_note_after_table(table, text: str) -> None:
    if next_paragraph_text(table).startswith("Chú thích:"):
        return
    paragraph_after_table(table, text, "DATN Note")


def add_or_update_figure_note(caption, text: str) -> None:
    node = caption._p.getnext()
    note = None
    if node is not None and node.tag == qn("w:p"):
        candidate = Paragraph(node, caption._parent)
        if candidate.text.strip().startswith("Chú giải chức năng:"):
            note = candidate
    if note is None:
        note = paragraph_after(caption, style="DATN Figure Note")
    else:
        clear_paragraph(note)
        note.style = "DATN Figure Note"
    lead = note.add_run("Chú giải chức năng: ")
    set_run_font(lead, 13, bold=True)
    detail = note.add_run(text)
    set_run_font(detail, 13)
    note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    note.paragraph_format.first_line_indent = Cm(0)
    note.paragraph_format.line_spacing = 1.2
    note.paragraph_format.space_before = Pt(3)
    note.paragraph_format.space_after = Pt(6)


def set_cell_text(cell, text: str) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            set_run_font(run, 11)


def mark_repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        properties.append(header)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def format_tables(document: Document) -> None:
    for table in document.tables:
        table.autofit = True
        if table.rows:
            mark_repeat_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            prevent_row_split(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                    paragraph.paragraph_format.line_spacing = 1.15
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(2)
                    if row_index == 0:
                        paragraph.paragraph_format.keep_with_next = True
                    for run in paragraph.runs:
                        set_run_font(run, 11, bold=True if row_index == 0 else None)


def has_drawing(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))


def format_paragraphs(document: Document, cover_limit: int) -> None:
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else ""

        if index < cover_limit:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.2
            paragraph.paragraph_format.space_after = Pt(6)
            size = 18 if index <= 1 else 13
            for run in paragraph.runs:
                set_run_font(run, size, bold=index <= 1)
            continue

        if style_name.startswith("Heading"):
            for run in paragraph.runs:
                set_run_font(run)
            continue

        if text.startswith("Bảng "):
            paragraph.style = "DATN Table Caption"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                set_run_font(run, 13, bold=True)
            continue

        if text.startswith("Hình "):
            paragraph.style = "DATN Figure Caption"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = False
            for run in paragraph.runs:
                set_run_font(run, 13, italic=True)
            continue

        if text.startswith("Chú thích:"):
            paragraph.style = "DATN Note"
            paragraph.paragraph_format.first_line_indent = Cm(0)
            for run in paragraph.runs:
                set_run_font(run, 13, italic=True)
            continue

        if text.startswith("Chú giải chức năng:"):
            paragraph.style = "DATN Figure Note"
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.2
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(6)
            for run in paragraph.runs:
                set_run_font(run, 13)
            continue

        if has_drawing(paragraph):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_with_next = True
            continue

        if style_name.startswith("TOC"):
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.15
            for run in paragraph.runs:
                set_run_font(run, 12)
            continue

        if style_name.startswith("List"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.line_spacing = 1.2
            for run in paragraph.runs:
                set_run_font(run, 13)
            continue

        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Cm(1)
        paragraph.paragraph_format.line_spacing = 1.2
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_run_font(run, 13)


def update_srs_content(document: Document) -> None:
    replacements = {
        "9. USE CASE": "9. CA SỬ DỤNG",
        "Phụ lục A Yêu cầu contract 0.5.0": "PHỤ LỤC C — YÊU CẦU CONTRACT 0.5.0",
        "A.1 Mục tiêu": "C.1 Mục tiêu",
        "A.2 Tiêu chí nghiệm thu": "C.2 Tiêu chí nghiệm thu",
        "Revision 1.2.2 Full transcript and full video quiz": "PHỤ LỤC D — FULL TRANSCRIPT VÀ FULL-VIDEO QUIZ",
        "A.1 Thay đổi luồng nghiệp vụ": "D.1 Thay đổi luồng nghiệp vụ",
        "A.2 Quy tắc số câu theo thời lượng": "D.2 Quy tắc số câu theo thời lượng",
    }
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value in replacements:
            set_text(paragraph, replacements[value])
        if paragraph.text.strip() == "9. CA SỬ DỤNG":
            paragraph.style = "Heading 1"

    appendix = None
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value.startswith("PHỤ LỤC C"):
            appendix = "C"
            paragraph.style = "Heading 1"
            continue
        if value.startswith("PHỤ LỤC D"):
            appendix = "D"
            paragraph.style = "Heading 1"
            continue
        if appendix == "C":
            for old, new in (
                ("A.1 ", "C.1 "),
                ("A.2 ", "C.2 "),
                ("A.3 ", "C.3 "),
            ):
                if value.startswith(old):
                    set_text(paragraph, new + value[len(old):])
                    paragraph.style = "Heading 2"
                    break
        elif appendix == "D":
            if value in {"Quy tắc hoàn tất và rollback", "D.1 Quy tắc hoàn tất và rollback"}:
                set_text(paragraph, "D.1 Quy tắc hoàn tất và rollback")
                paragraph.style = "Heading 2"
            elif value in {"Số câu mục tiêu", "D.2 Số câu mục tiêu"}:
                set_text(paragraph, "D.2 Số câu mục tiêu")
                paragraph.style = "Heading 2"

    if len(document.tables) != 15:
        raise RuntimeError(f"SRS table count changed: expected 15, got {len(document.tables)}")

    headers = document.tables[7].rows[0].cells
    if len(headers) >= 3:
        set_cell_text(headers[1], "Ca sử dụng")
        set_cell_text(headers[2], "Tác nhân")

    captions = [
        "Bảng 0.1: Lịch sử phiên bản",
        "Bảng 1.1: Thuật ngữ và định nghĩa",
        "Bảng 4.1: Tác nhân và vai trò",
        "Bảng 5.1: Yêu cầu chức năng",
        "Bảng 5.2: Quy tắc trạng thái kích hoạt toàn cục",
        "Bảng 6.1: Yêu cầu phi chức năng",
        "Bảng 8.1: Mô hình dữ liệu nghiệp vụ",
        "Bảng 9.1: Danh sách ca sử dụng",
        "Bảng 10.1: Quy tắc nghiệp vụ",
        "Bảng 11.1: Công nghệ và trách nhiệm tích hợp",
        "Bảng 12.1: Xử lý ngoại lệ",
        "Bảng 14.1: Yêu cầu mở rộng sau MVP",
        "Bảng A.1: Cấu hình MVP đề xuất",
        "Bảng C.1: Tiêu chí nghiệm thu contract 0.5.0",
        "Bảng D.1: Số câu mục tiêu theo thời lượng",
    ]
    for table, caption in zip(document.tables, captions, strict=True):
        add_table_caption(table, caption)

    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("Bảng 9.1:"):
            paragraph.paragraph_format.page_break_before = False

    add_note_after_table(
        document.tables[9],
        "Chú thích: Extension là tiện ích trình duyệt; Backend là dịch vụ nghiệp vụ trung tâm; "
        "AI Service là dịch vụ FastAPI điều phối AI; LLM Service là dịch vụ mô hình ngôn ngữ; "
        "Database là kho dữ liệu PostgreSQL.",
    )
    add_note_after_table(
        document.tables[12],
        "Chú thích: Activation là trạng thái bật/tắt StudyLens; Question Count là số câu mục tiêu; "
        "Question Type là loại câu hỏi; AI Serving là cách cung cấp dịch vụ AI; Database là hệ quản trị dữ liệu.",
    )


def render_dot(source: str) -> bytes:
    result = subprocess.run(
        ["dot", "-Tpng", "-Gdpi=180"],
        input=source.encode("utf-8"),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=True,
    )
    return result.stdout


def add_diagram_after(anchor, heading: str, description: str, source: str, caption: str):
    heading_paragraph = paragraph_after(anchor, heading, "Heading 3")
    description_paragraph = paragraph_after(heading_paragraph, description, "Normal")
    image_paragraph = paragraph_after(description_paragraph)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.first_line_indent = Cm(0)
    run = image_paragraph.add_run()
    run.add_picture(BytesIO(render_dot(source)), width=Cm(15.2))
    caption_paragraph = paragraph_after(image_paragraph, caption, "DATN Figure Caption")
    return caption_paragraph


def replace_all_sds_diagrams(document: Document) -> None:
    module_path = ROOT / "scripts" / "update-week6-docs.py"
    spec = importlib.util.spec_from_file_location("week6_docs", module_path)
    if spec is None or spec.loader is None:
        raise RuntimeError("Cannot load the existing diagram definitions.")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    sources = [
        module.DIAGRAMS[0],
        LOGICAL_DOT,
        PHYSICAL_DOT,
        *module.DIAGRAMS[1:],
    ]
    if len(document.inline_shapes) != len(sources):
        raise RuntimeError(
            f"SDS diagram count changed: expected {len(sources)}, got {len(document.inline_shapes)}"
        )
    for shape, source in zip(document.inline_shapes, sources, strict=True):
        relationship_id, _ = document.part.get_or_add_image(BytesIO(render_dot(source)))
        shape._inline.graphic.graphicData.pic.blipFill.blip.embed = relationship_id


def update_sds_content(document: Document) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "Quản lý tài liệu":
            set_text(paragraph, "Bảng 0.1: Thông tin quản lý tài liệu")
            paragraph.style = "DATN Table Caption"
        elif paragraph.text.strip() == "Lịch sử phiên bản":
            set_text(paragraph, "LỊCH SỬ PHIÊN BẢN")
            paragraph.style = "Heading 1"

    if len(document.tables) != 9:
        raise RuntimeError(f"SDS table count changed: expected 9, got {len(document.tables)}")

    add_table_caption(document.tables[1], "Bảng 0.2: Lịch sử phiên bản")
    headers = document.tables[5].rows[0].cells
    if len(headers) >= 2:
        set_cell_text(headers[0], "Phương thức")
        set_cell_text(headers[1], "Điểm cuối (endpoint)")

    add_note_after_table(
        document.tables[8],
        "Chú thích: Video Context Reader đọc ngữ cảnh video; Persistent Activation Manager quản lý trạng thái bật bền vững; "
        "Activation Store lưu trạng thái kích hoạt; Session Manager quản lý phiên; Video Change Handler xử lý đổi video; "
        "Transcript Reader/Processor thu thập và chuẩn hóa transcript; LearningPackageService cung cấp gói học tập; "
        "Question Generator sinh câu hỏi; Grading Service chấm bài; Player Controller điều khiển tua video; "
        "History API/UI cung cấp lịch sử; Error Handler chuẩn hóa xử lý lỗi.",
    )

    if not any(p.text.strip().startswith("Hình 3.2:") for p in document.paragraphs):
        heading_32 = next(
            paragraph for paragraph in document.paragraphs
            if paragraph.text.strip().startswith("3.2 ")
        )
        anchor = heading_32._p.getprevious()
        while anchor is not None and anchor.tag != qn("w:p"):
            anchor = anchor.getprevious()
        if anchor is None:
            raise RuntimeError("Cannot locate the insertion point before section 3.2.")
        anchor_paragraph = Paragraph(anchor, heading_32._parent)
        last = add_diagram_after(
            anchor_paragraph,
            "3.1.1 Logical Architecture",
            "Kiến trúc logic phân tách hệ thống theo các lớp trình bày, ứng dụng, tích hợp và dữ liệu/dịch vụ ngoài. "
            "Sơ đồ thể hiện ranh giới trách nhiệm và hướng trao đổi dữ liệu, không phụ thuộc vào môi trường triển khai vật lý.",
            LOGICAL_DOT,
            "Hình 3.2: Logical Architecture của StudyLens AI",
        )
        add_diagram_after(
            last,
            "3.1.2 Physical Architecture",
            "Kiến trúc vật lý mô tả các tiến trình và container khi triển khai cục bộ. Chỉ API Backend công bố trên loopback; "
            "AI Service và PostgreSQL nằm trong mạng Docker nội bộ, còn dữ liệu được duy trì trong named volume.",
            PHYSICAL_DOT,
            "Hình 3.3: Physical Architecture của StudyLens AI",
        )

    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs[:-1]):
        if paragraph.text.strip() in {
            "3.1.1 Logical Architecture",
            "3.1.2 Physical Architecture",
        }:
            paragraph.paragraph_format.keep_with_next = True
            paragraphs[index + 1].paragraph_format.keep_with_next = True

        if paragraph.text.strip() == "3.1.2 Physical Architecture":
            set_text(
                paragraphs[index + 1],
                "Kiến trúc vật lý mô tả các tiến trình và container được triển khai trên StudyLens Server chạy Docker. "
                "API Backend là điểm truy cập của Extension; AI Service và PostgreSQL chỉ giao tiếp trong mạng Docker nội bộ, "
                "còn dữ liệu bền vững được duy trì trong named volume.",
            )

    figure_notes = {
        "Hình 3.1:": (
            "StudyLens Extension đọc ngữ cảnh YouTube, thu thập full transcript và hiển thị Side Panel; "
            "ASP.NET Core Backend quản lý phiên, durable job, quiz, chấm điểm và lịch sử; FastAPI AI Service xây prompt và kiểm tra structured output; "
            "PostgreSQL lưu dữ liệu nghiệp vụ; Google Gemini API sinh transcript dự phòng, câu hỏi và kết quả AI."
        ),
        "Hình 3.2:": (
            "Presentation Layer cung cấp Side Panel và bộ đọc ngữ cảnh YouTube; Application Layer điều phối kích hoạt, phiên học, quiz, chấm điểm và lịch sử; "
            "Integration Layer cung cấp REST API và cổng FastAPI; Data & External Services lưu dữ liệu và kết nối YouTube Timedtext cùng Google Gemini API."
        ),
        "Hình 3.3:": (
            "User Device chạy Chrome/Edge và StudyLens Extension; StudyLens Server (Docker Host) chứa Migration Job, API Container, AI Container, PostgreSQL 16 và Named Data Volume; "
            "Migration Job nâng cấp schema trước khi API khởi động; API Container xử lý nghiệp vụ; AI Container gọi Gemini; PostgreSQL và Named Data Volume lưu trữ bền vững."
        ),
        "Hình 4.1:": (
            "Content Script đọc video context và caption JSON3/XML; Service Worker cô lập active tab, duy trì trạng thái và gọi Backend; Side Panel hiển thị tiến trình, transcript, quiz và lịch sử; "
            "chrome.storage.local lưu ON/OFF cùng preferences; ASP.NET Core Backend là đích REST duy nhất của Extension."
        ),
        "Hình 5.1:": (
            "HTTP Endpoints tiếp nhận session, transcript, attempt và history; Application Services xác thực dữ liệu, bảo đảm idempotency và transaction; "
            "Background Worker chạy durable job, retry và recovery; AI Clients gọi FastAPI cho transcript fallback, quiz và grading; Entity Framework Core thực hiện persistence và migration với PostgreSQL."
        ),
        "Hình 6.1:": (
            "Backend Request mang videoId, full transcript hoặc short answer; FastAPI Feature Service chọn luồng transcript, question generation hay grading; Prompt Builder tạo ngữ cảnh và evidence; "
            "Gemini Provider gọi API bằng key chỉ có trên server; Pydantic Validation kiểm tra schema, timestamp và lỗi an toàn; Backend Response trả transcript, quiz hoặc kết quả chấm."
        ),
        "Hình 9.1:": (
            "Bước 1 tạo phiên và lấy full caption; bước 2 xác thực transcript và lưu durable job; bước 3 sinh transcript dự phòng hoặc quiz; bước 4 công bố quiz hợp lệ; "
            "bước 5 gửi toàn bộ bài làm; bước 6 chấm atomically, lưu history và trả evidence timestamp."
        ),
        "Hình 11.1:": (
            "Study, History và Settings là ba khu vực giao diện; Video title + active session xác định video/phiên hiện tại; Processing hiển thị tiến trình; Transcript ready xác nhận nguồn transcript; "
            "Full-video quiz hiển thị toàn bộ câu hỏi và nút nộp; Score + explanation hiển thị điểm, giải thích và mốc xem lại."
        ),
        "Hình 16.1:": (
            "api container công bố Backend cho Extension; migrate one-shot job chạy EF Core migration trước API; ai container chạy FastAPI trong mạng Docker nội bộ; "
            "postgres container lưu dữ liệu trong named volume; Google Gemini API là dịch vụ ngoài được AI Service gọi qua HTTPS."
        ),
    }
    for paragraph in list(document.paragraphs):
        caption_text = paragraph.text.strip()
        for prefix, note in figure_notes.items():
            if caption_text.startswith(prefix):
                add_or_update_figure_note(paragraph, note)
                break

    replace_all_sds_diagrams(document)


def format_report(path: Path, *, kind: str) -> None:
    document = Document(path)
    configure_styles(document)
    if kind == "srs":
        update_srs_content(document)
        cover_limit = 5
    else:
        update_sds_content(document)
        cover_limit = 6
    configure_sections(document)
    format_tables(document)
    format_paragraphs(document, cover_limit)
    document.core_properties.title = (
        "StudyLens AI - Software Requirements Specification"
        if kind == "srs"
        else "StudyLens AI - Software Design Specification"
    )
    document.save(path)


def main() -> None:
    if not SRS_PATH.exists() or not SDS_PATH.exists():
        raise FileNotFoundError("The SRS or SDS source document is missing.")
    format_report(SRS_PATH, kind="srs")
    format_report(SDS_PATH, kind="sds")
    print("Formatted both StudyLens AI thesis reports.")


if __name__ == "__main__":
    main()

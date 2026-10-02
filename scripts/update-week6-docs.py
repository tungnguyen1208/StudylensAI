from __future__ import annotations

from collections import Counter
from datetime import date
from pathlib import Path
import subprocess

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SRS_PATH = ROOT / "SRS_StudyLens_AI_Persistent_Activation_v1.2 (2).docx"
SDS_PATH = ROOT / "SDS - StudyLens AI - Học tập có tương tác với Youtube.docx"
REVISION_DATE = date(2026, 10, 2).strftime("%d/%m/%Y")


DIAGRAMS = [
    r'''
digraph overview {
  graph [bgcolor="white", pad="0.2", ranksep="0.55", nodesep="0.45"];
  node [shape=box, style="rounded,filled", fillcolor="#EEF5FB", color="#496B8A", fontname="Arial", fontsize=13, margin="0.18,0.12"];
  edge [color="#496B8A", fontcolor="#496B8A", fontname="Arial", fontsize=10, penwidth=1.4];
  ext [label="StudyLens Extension\nYouTube context + full captions\nSide Panel"];
  api [label="ASP.NET Core Backend\nSessions + durable jobs\nQuiz + grading + history"];
  ai [label="FastAPI AI Service\nPrompt orchestration\nStructured output validation"];
  gemini [shape=ellipse, fillcolor="#FFF6E5", color="#9B6E24", label="Google AI Gemini API\nTranscript fallback + inference"];
  db [shape=cylinder, fillcolor="#EFF7EC", color="#48765B", label="PostgreSQL 16\nBusiness source of truth"];
  ext -> api [label="contract 0.5.0"];
  api -> ai [label="internal AI API"];
  ai -> gemini [label="HTTPS"];
  api -> db [label="EF Core"];
}
''',
    r'''
digraph extension {
  graph [bgcolor="white", pad="0.2", ranksep="0.45", nodesep="0.35"];
  node [shape=box, style="rounded,filled", fillcolor="#EEF5FB", color="#496B8A", fontname="Arial", fontsize=12, margin="0.15,0.1"];
  edge [color="#496B8A", fontcolor="#496B8A", fontname="Arial", fontsize=9, penwidth=1.3];
  yt [label="YouTube Web\nPlayer + captionTracks"];
  cs [label="Content Script\nVideo context + JSON3/XML captions\nDOM fallback guard"];
  sw [label="Service Worker\nActive-tab isolation + Backend API"];
  panel [label="Side Panel\nProcessing + transcript + quiz + history"];
  storage [shape=cylinder, fillcolor="#EFF7EC", color="#48765B", label="chrome.storage.local\nON/OFF + preferences"];
  api [label="ASP.NET Core Backend"];
  yt -> cs [label="context / full cues"];
  cs -> sw [label="versioned messages"];
  sw -> api [label="REST 0.5.0"];
  sw -> panel [dir=both, label="state / actions"];
  sw -> storage [dir=both, label="persistent gate"];
}
''',
    r'''
digraph backend {
  graph [bgcolor="white", pad="0.2", ranksep="0.5", nodesep="0.35"];
  node [shape=box, style="rounded,filled", fillcolor="#EEF5FB", color="#496B8A", fontname="Arial", fontsize=12, margin="0.15,0.1"];
  edge [color="#496B8A", fontcolor="#496B8A", fontname="Arial", fontsize=9, penwidth=1.3];
  ext [label="StudyLens Extension"];
  endpoints [label="HTTP Endpoints\nSessions + transcript + attempts + history"];
  services [label="Application Services\nValidation + idempotency + transactions"];
  jobs [label="Background Worker\nDurable ProcessingJobs + retry/recovery"];
  ai [label="AI Clients\nTranscript fallback + quiz + grading"];
  ef [label="Entity Framework Core\nExplicit migrations"];
  db [shape=cylinder, fillcolor="#EFF7EC", color="#48765B", label="PostgreSQL 16"];
  fast [label="FastAPI AI Service"];
  ext -> endpoints [label="REST"];
  endpoints -> services;
  services -> jobs;
  services -> ef;
  jobs -> ai;
  ai -> fast [label="HTTP"];
  jobs -> ef;
  ef -> db;
}
''',
    r'''
digraph ai {
  graph [bgcolor="white", pad="0.2", ranksep="0.45", nodesep="0.35"];
  node [shape=box, style="rounded,filled", fillcolor="#EEF5FB", color="#496B8A", fontname="Arial", fontsize=12, margin="0.15,0.1"];
  edge [color="#496B8A", fontcolor="#496B8A", fontname="Arial", fontsize=9, penwidth=1.3];
  req [label="Backend Request\nvideoId fallback | full transcript | short answer"];
  route [label="FastAPI Feature Service\nTranscript | Question Generation | Grading"];
  prompt [label="Prompt Builder\nFull-video context + preferences + evidence"];
  provider [label="Gemini Provider\nServer-only API key + timeout mapping"];
  schema [label="Pydantic Validation\nSchema + timestamps + safe errors"];
  res [label="Backend Response\nTimestamped transcript | quiz | grade"];
  gemini [shape=ellipse, fillcolor="#FFF6E5", color="#9B6E24", label="Google AI Gemini API"];
  req -> route -> prompt -> provider;
  provider -> gemini [label="HTTPS"];
  gemini -> provider;
  provider -> schema -> res;
}
''',
    r'''
digraph sequence {
  graph [bgcolor="white", pad="0.2", ranksep="0.35", nodesep="0.25"];
  node [shape=box, style="rounded,filled", fillcolor="#EEF5FB", color="#496B8A", fontname="Arial", fontsize=11, margin="0.16,0.1"];
  edge [color="#496B8A", fontcolor="#3F5F7A", fontname="Arial", fontsize=9, penwidth=1.3];
  acquire [label="1. Extension\nCreate session + acquire full caption"];
  persist [label="2. Backend + PostgreSQL\nValidate transcript + persist durable job", fillcolor="#EFF7EC", color="#48765B"];
  generate [label="3. FastAPI + Gemini\nTranscript fallback or full-video quiz", fillcolor="#FFF6E5", color="#9B6E24"];
  publish [label="4. Backend + PostgreSQL\nValidate DTO + publish ready quiz", fillcolor="#EFF7EC", color="#48765B"];
  attempt [label="5. Extension\nPoll package + submit full attempt"];
  result [label="6. Backend / FastAPI\nGrade atomically + save history + return evidence"];
  acquire -> persist -> generate -> publish -> attempt -> result;
}
''',
    r'''
digraph panel {
  graph [bgcolor="white", pad="0.2", ranksep="0.25", nodesep="0.2"];
  node [shape=box, style="rounded,filled", fillcolor="#F8FBF4", color="#7AA35A", fontname="Arial", fontsize=11, margin="0.16,0.1"];
  title [label="StudyLens AI\nStudy | History | Settings", fillcolor="#E9F5DF", fontsize=14];
  video [label="Video title + active session"];
  progress [label="Processing 75%\nAI Service is generating the full-video quiz"];
  transcript [label="Transcript ready\nYouTube captions or Gemini video source"];
  quiz [label="Full-video quiz\nQuestion 1/5 ... Question 5/5\nSubmit full attempt"];
  result [label="Score + explanation\nReview evidence timestamp"];
  title -> video -> progress -> transcript -> quiz -> result;
}
''',
    r'''
digraph deployment {
  graph [bgcolor="white", pad="0.2", ranksep="0.5", nodesep="0.35"];
  node [shape=box, style="rounded,filled", fillcolor="#EEF5FB", color="#496B8A", fontname="Arial", fontsize=12, margin="0.16,0.1"];
  edge [color="#496B8A", fontcolor="#496B8A", fontname="Arial", fontsize=9, penwidth=1.3];
  ext [label="Chrome / Edge\nStudyLens Extension"];
  api [label="api container\nASP.NET Core 8\n127.0.0.1:5000 only"];
  migrate [label="migrate one-shot job\n--migrate-only"];
  ai [label="ai container\nPython 3.12 FastAPI\nprivate Docker network"];
  db [shape=cylinder, fillcolor="#EFF7EC", color="#48765B", label="postgres container\nPostgreSQL 16 + named volume\nprivate Docker network"];
  gem [shape=ellipse, fillcolor="#FFF6E5", color="#9B6E24", label="Google AI Gemini API"];
  ext -> api [label="HTTP / REST"];
  migrate -> db [label="EF Core migrations"];
  api -> db [label="EF Core"];
  api -> ai [label="HTTP"];
  ai -> gem [label="HTTPS"];
}
''',
]


def replace_diagrams(document: Document) -> None:
    if len(document.inline_shapes) < len(DIAGRAMS):
        raise RuntimeError("The SDS does not contain the expected diagram placeholders.")
    for index, source in enumerate(DIAGRAMS):
        rendered = subprocess.run(
            ["dot", "-Tpng", "-Gdpi=170"],
            input=source.encode("utf-8"),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=True,
        ).stdout
        shape = document.inline_shapes[index]
        relationship_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        document.part.related_parts[relationship_id]._blob = rendered


def set_text(paragraph, text: str) -> None:
    paragraph.text = text
    for run in paragraph.runs:
        run.font.name = "Arial"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")


def set_cell(cell, text: str) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")


def set_row(table, row_index: int, values: list[str]) -> None:
    for cell, value in zip(table.rows[row_index].cells, values, strict=True):
        set_cell(cell, value)


def set_table_widths(table, widths: list[float]) -> None:
    table.autofit = False
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Inches(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = Inches(width)


def set_table_font_size(table, size: float) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(size)


def add_history_row(table, values: list[str]) -> None:
    for row_index, existing_row in enumerate(table.rows):
        if existing_row.cells[0].text.strip() == values[0]:
            set_row(table, row_index, values)
            for cell in existing_row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
            return
    row = table.add_row()
    for cell, value in zip(row.cells, values, strict=True):
        set_cell(cell, value)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(10)


def repeat_table_header(table) -> None:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_tables(document: Document) -> None:
    for table in document.tables:
        explicit_sizes = Counter(
            run.font.size.pt
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
            for run in paragraph.runs
            if run.font.size is not None
        )
        table_font_size = Pt(explicit_sizes.most_common(1)[0][0] if explicit_sizes else 9)
        repeat_table_header(table)
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tc_pr = cell._tc.get_or_add_tcPr()
                margins = tc_pr.first_child_found_in("w:tcMar")
                if margins is None:
                    margins = OxmlElement("w:tcMar")
                    tc_pr.append(margins)
                for edge in ("top", "start", "bottom", "end"):
                    node = margins.find(qn(f"w:{edge}"))
                    if node is None:
                        node = OxmlElement(f"w:{edge}")
                        margins.append(node)
                    node.set(qn("w:w"), "100")
                    node.set(qn("w:type"), "dxa")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
                        if run.font.size is None:
                            run.font.size = table_font_size


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text, style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)


def add_revision_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_row(table, 0, headers)
    for row_values in rows:
        row = table.add_row()
        for cell, value in zip(row.cells, row_values, strict=True):
            set_cell(cell, value)
    repeat_table_header(table)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths, strict=True):
                cell.width = width


def update_srs() -> None:
    document = Document(SRS_PATH)
    paragraphs = document.paragraphs

    set_text(paragraphs[3], "\nPhiên bản: 1.2.2\nNgày: 02/10/2026\nPhạm vi: MVP contract 0.5.0")
    set_text(paragraphs[11], (
        "StudyLens AI hoạt động trên YouTube Web với công tắc ON/OFF toàn cục do người dùng điều khiển. "
        "Khi ON và trang xem hợp lệ, Backend tạo Study Session ngay; Extension lấy toàn bộ caption có timestamp theo contract 0.5.0 "
        "qua captionTracks, Timedtext JSON3 rồi XML. DOM chỉ là bản xem trước khi chưa chứng minh đủ video. "
        "Nếu caption không khả dụng hoặc không đủ, Backend yêu cầu FastAPI dùng Gemini phân tích URL video công khai. "
        "Extension không gửi audio và không chứa API key."
    ))
    set_text(paragraphs[31], "Mỗi video có một quiz toàn video; số câu mục tiêu là 3/5/10/15 theo các mốc thời lượng 10/30/60 phút.")
    set_text(paragraphs[49], "Hiển thị rõ trạng thái StudyLens OFF, ON, đang xử lý, sẵn sàng hoặc lỗi có thể thử lại.")
    set_text(paragraphs[50], "Hiển thị tiến trình theo bước: tạo phiên, lấy transcript, xác thực, Gemini fallback, tạo quiz và quiz sẵn sàng.")
    set_text(paragraphs[51], "Hiển thị transcript YouTube trực tiếp tách biệt với trạng thái transcript đã được Backend xác thực.")
    set_text(paragraphs[52], "Hiển thị toàn bộ câu hỏi trong một quiz và không lộ đáp án trước khi nộp cả bài.")
    set_text(paragraphs[53], "Sau khi chấm, hiển thị điểm, đáp án tham chiếu, giải thích và nút xem lại timestamp hợp lệ.")
    set_text(paragraphs[56], "Số câu mục tiêu: <10 phút là 3; 10-<30 phút là 5; 30-<60 phút là 10; >=60 phút là 15.")
    set_text(paragraphs[57], "Question Type: Multiple Choice / Short Answer.")
    set_text(paragraphs[58], "Difficulty: Easy / Medium / Hard.")
    set_text(paragraphs[64], "PostgreSQL là cơ sở dữ liệu bền vững của bản triển khai. Migration Entity Framework Core phải chạy xong trước Backend; restart không được làm mất session, job, quiz, attempt hoặc history.")

    main_flow = [
        "1. Người dùng mở video YouTube hợp lệ trong khi StudyLens đang ON.",
        "2. Extension đọc trạng thái ON đã lưu và ngữ cảnh video hiện tại.",
        "3. Backend tạo Study Session active ngay, gắn với youtubeVideoId, title và preference snapshot.",
        "4. Extension lấy full transcript có timestamp ngay từ caption track, không chờ video phát hết.",
        "5. Extension chuẩn hóa, hash và gửi transcript một lần; Backend xác thực và lưu capture/cue.",
        "6. Nếu caption không đủ, Backend xếp job transcriptGenerate để FastAPI gửi URL video công khai cho Gemini.",
        "7. Khi transcript ready, Backend xếp một job quizGenerate cho toàn video.",
        "8. FastAPI sinh tối đa số câu mục tiêu theo thời lượng, kiểm tra schema và timestamp bằng chứng.",
        "9. Side Panel polling learning package và chỉ công bố quiz khi trạng thái ready.",
        "10. Người học gửi toàn bộ câu trả lời trong một request; Backend/FastAPI chấm atomically.",
        "11. Backend chỉ đánh dấu session completed sau khi attempt hoàn chỉnh được lưu.",
        "12. History gắn với video, session, quiz attempt và evidence timestamp; người học có thể xem lại đúng video.",
    ]
    for index, text in enumerate(main_flow, start=67):
        set_text(paragraphs[index], text)

    set_text(paragraphs[82], "Kiến trúc logic gồm năm thành phần: Chrome/Edge Extension, ASP.NET Core Backend, FastAPI AI Service, Google AI Gemini API và PostgreSQL. Extension chỉ gọi Backend; Backend là nguồn sự thật; FastAPI là biên provider duy nhất gọi Gemini.")
    data_flow = [
        "Khi trạng thái đã lưu là ON, Extension xác nhận trang xem YouTube và tạo session ngay qua Backend.",
        "Extension không gửi yêu cầu phân loại video học tập/giải trí.",
        "Extension lấy full caption qua Timedtext JSON3/XML, chuẩn hóa cue và gửi cue-only JSON một lần; không gửi audio.",
        "Backend xác thực hash, lưu transcript và xếp ProcessingJob bền vững.",
        "Khi caption thiếu, FastAPI nhận videoId từ Backend, tự tạo URL YouTube công khai và yêu cầu Gemini trả transcript timestamp.",
        "Khi transcript ready, Backend gửi full transcript và preference snapshot sang FastAPI để sinh quiz toàn video.",
        "Backend chỉ trả quiz public không có answer key; sau batch submit mới trả reference answer, explanation và evidence.",
        "PostgreSQL lưu video, session, transcript, job, quiz, attempt và history; idempotency ngăn dữ liệu trùng khi retry.",
    ]
    for index, text in enumerate(data_flow, start=84):
        set_text(paragraphs[index], text)

    acceptance = [
        "1. Extension cài được trên Chrome/Edge Manifest V3 và có ID ổn định trong gói release.",
        "2. ON/OFF được lưu và khôi phục sau khi chuyển video hoặc khởi động lại trình duyệt.",
        "3. Khi ON, mở video hợp lệ tạo session active ngay và lấy full transcript mà không chờ xem hết.",
        "4. Khi OFF, rời /watch hoặc chuyển A sang B, response cũ không được cập nhật video mới.",
        "5. Transcript YouTube là nguồn ưu tiên; video công khai thiếu caption được thử Gemini URL fallback.",
        "6. Quiz là một bài cho toàn video với mục tiêu 3/5/10/15 câu theo các mốc 10/30/60 phút.",
        "7. Quiz public không lộ đáp án; batch attempt được chấm và lưu atomically.",
        "8. Lỗi transcript, sinh quiz hoặc grading không đánh dấu session completed và không công bố dữ liệu bán phần.",
        "9. Retry đúng bước và replay cùng idempotency key không tạo capture, quiz hay attempt trùng.",
        "10. History tồn tại sau restart, gắn đúng video/session và cho phép seek evidence hợp lệ.",
        "11. PostgreSQL migration, health readiness, backup/restore và job recovery phải đạt trước release.",
        "12. API key chỉ ở server; PostgreSQL/FastAPI không publish port host trong Docker release.",
        "13. Automated suites cho Extension, Backend, AI Service và contract phải pass.",
        "14. Chrome và Edge phải pass ma trận M-01…M-08 với screenshot/log đã làm sạch.",
        "15. Ba Gemini smoke case phải đạt hoặc được ghi rõ là bị chặn bởi quota/provider; không đánh dấu release hoàn tất khi còn bị chặn.",
    ]
    for index, text in enumerate(acceptance, start=94):
        set_text(paragraphs[index], text)

    set_text(paragraphs[112], "Phiên bản 1.2.2 thay thế luồng contract 0.4.0 theo chu kỳ bằng contract 0.5.0: session tạo ngay, full transcript toàn video, Gemini URL fallback, một quiz toàn video, batch attempt, history theo video và PostgreSQL bền vững.")
    set_text(paragraphs[115], "Phụ lục A Yêu cầu contract 0.5.0")
    set_text(paragraphs[116], "Phụ lục này làm rõ FR-03…FR-15 cho luồng full transcript và full-video quiz. Khi có khác biệt với nội dung cũ, revision 1.2.2 và contract 0.5.0 được ưu tiên.")
    set_text(paragraphs[118], "Nguồn 1 là captionTracks trong trang xem YouTube; Extension tải Timedtext JSON3, sau đó XML. DOM chỉ được gửi nếu chứng minh bao phủ gần toàn bộ thời lượng. Nguồn 2 là Gemini phân tích URL video công khai khi caption thiếu; Extension không gửi audio.")
    set_text(paragraphs[120], "Cue được chuẩn hóa, sắp xếp và hash trước khi gửi. Backend xác thực, lưu transcript và xếp đúng một job quiz cho toàn video. FastAPI nhận full transcript cùng questionType/difficulty; không có timer, interval hoặc transcript segment định kỳ.")

    tables = document.tables
    add_history_row(tables[0], ["1.2.2", REVISION_DATE, "Nhóm StudyLens AI", "Contract 0.5.0, full transcript/full-video quiz, Gemini URL fallback, batch attempt, PostgreSQL và release hardening."])
    set_row(tables[1], 3, ["Transcript", "Bản chép lời toàn video gồm cue {startMs, endMs, text}; nguồn youtubeCaption hoặc geminiVideo đã được Backend xác thực."])
    set_row(tables[1], 5, ["Quiz", "Một bài kiểm tra cho toàn video; số câu mục tiêu phụ thuộc thời lượng."])
    set_row(tables[1], 7, ["AI Service", "FastAPI nhận yêu cầu từ Backend để tạo transcript fallback, sinh quiz, chấm và giải thích; chỉ thành phần này gọi Gemini."])
    set_row(tables[2], 3, ["StudyLens Extension", "Lưu ON/OFF, đọc video context, lấy full caption, hiển thị trạng thái/quiz/history và chỉ gọi Backend."])
    set_row(tables[2], 4, ["ASP.NET Core Backend", "Nguồn sự thật cho video, session, transcript, job, quiz, attempt và history; xử lý idempotency/retry."])
    set_row(tables[2], 5, ["FastAPI AI Service", "Nhận full transcript hoặc videoId fallback từ Backend; gọi Gemini, kiểm tra structured output và không lưu API key phía client."])

    functional_rows = {
        3: ["FR-03", "Quản lý Study Session", "Khi ON và video hợp lệ, Backend SHALL tạo session active ngay, gắn youtubeVideoId/title/preferences; chuyển video SHALL đóng context cũ và tạo session mới mà không tắt StudyLens."],
        4: ["FR-04", "Nhận và xác thực full transcript", "Extension SHALL lấy toàn bộ caption timestamp ngay khi mở video theo JSON3→XML; chuẩn hóa/hash và gửi một payload. Caption thiếu SHALL kích hoạt Gemini URL fallback cho video công khai."],
        5: ["FR-05", "Theo dõi trạng thái xử lý", "Side Panel SHALL khôi phục và hiển thị trạng thái transcript/quiz bền vững từ Backend; transcript xem trước SHALL phân biệt với bản ready."],
        6: ["FR-06", "Tạo learning package toàn video", "Backend SHALL dùng transcript ready cho một quiz toàn video; không SHALL tạo segment hoặc quiz theo chu kỳ xem."],
        7: ["FR-07", "Sinh câu hỏi bằng AI", "Backend SHALL gửi full transcript và preference snapshot tới FastAPI. Structured output SHALL đúng schema, bám evidence và có timestamp hợp lệ."],
        8: ["FR-08", "Hiển thị quiz", "Extension SHALL hiển thị toàn bộ câu hỏi trong Side Panel, hỗ trợ bàn phím và không lộ đáp án trước submit."],
        9: ["FR-09", "Nộp toàn bộ bài", "Hệ thống SHALL nhận toàn bộ câu trả lời trong một request idempotent và lưu attempt/answers atomically."],
        10: ["FR-10", "Chấm và giải thích", "MCQ SHALL chấm xác định; short answer SHALL qua FastAPI/Gemini. Chỉ sau submit mới trả reference answer, score và explanation."],
        11: ["FR-11", "Gợi ý xem lại", "Kết quả SHALL mang evidenceStartMs/evidenceEndMs thuộc đúng video và chỉ seek khi giá trị hợp lệ."],
        12: ["FR-12", "Lưu kết quả", "PostgreSQL SHALL lưu video, session completed, quiz attempt và answers; phiên lỗi không được coi là completed."],
        13: ["FR-13", "Xem lịch sử", "History SHALL gắn theo youtubeVideoId, sessionId và quizAttemptId, giới hạn 100 item mới nhất mà không cắt rời một attempt."],
        14: ["FR-14", "Cấu hình StudyLens", "Người dùng SHALL cấu hình loại câu hỏi và độ khó; hệ thống không còn cấu hình interval."],
        15: ["FR-15", "Xử lý lỗi và retry", "Mỗi bước lỗi SHALL giữ error state tối thiểu để retry đúng bước, không công bố quiz/attempt bán phần và không tạo dữ liệu trùng."],
    }
    for row_index, values in functional_rows.items():
        set_row(tables[3], row_index, values)

    set_row(tables[4], 2, ["Người dùng chuyển ON", "Lưu ON; nếu đang ở video hợp lệ thì tạo session ngay và bắt đầu lấy full transcript", "ON"])
    set_row(tables[4], 3, ["Mở hoặc chuyển video mới khi đang ON", "Đóng context cũ, hủy tác vụ stale, tạo session mới và lấy full transcript", "ON"])
    set_row(tables[4], 6, ["Người dùng chuyển OFF", "Lưu OFF; hủy acquisition/polling và đóng session chưa hoàn tất", "OFF"])
    set_row(tables[5], 5, ["NFR-05", "Khả năng bảo trì", "Extension, Backend, FastAPI AI Service, Gemini provider và PostgreSQL SHALL có ranh giới module rõ ràng; contract 0.5.0 là seam tích hợp."])
    set_row(tables[5], 10, ["NFR-10", "Triển khai", "Docker release SHALL chạy non-root, chỉ publish Backend trên loopback, chạy migration trước web và không đưa secret vào image/bundle."])
    data_rows = [
        ["User preferences", "QuestionType, Difficulty", "Preference snapshot cho session; ON/OFF được lưu trong chrome.storage.local."],
        ["Video", "YouTube video ID, URL, title, duration", "Danh tính video dùng chung cho session, transcript, quiz và history."],
        ["Transcript", "Source, language, content hash, status, timestamp cues", "Bản transcript toàn video đã xác thực."],
        ["Study Session", "Video, preferences, transcript/quiz/session status", "Nguồn trạng thái cho learning package; chỉ completed sau full attempt."],
        ["Processing Job", "Type, idempotency key, attempts, error", "Job bền vững cho transcriptGenerate/quizGenerate và recovery sau restart."],
        ["Quiz", "Questions, options, evidence, model, prompt version", "Một quiz cho toàn video; answer key không công khai trước submit."],
        ["Quiz Attempt", "Session, answers, score, feedback, submitted time", "Batch attempt atomically, idempotent và gắn lịch sử theo video."],
        ["Learning Result", "Outcome, explanation, evidence timestamp", "Kết quả sau submit và điểm xem lại video."],
    ]
    for row_index, values in enumerate(data_rows, start=1):
        set_row(tables[6], row_index, values)

    set_row(tables[7], 4, ["UC-04", "Lấy full transcript", "Extension/Backend", "FR-04, FR-05", "Transcript YouTube ready hoặc Gemini fallback được xếp job."])
    set_row(tables[7], 5, ["UC-05", "Sinh quiz toàn video", "Backend/AI Service", "FR-06, FR-07", "Quiz ready với số câu mục tiêu theo duration."])
    set_row(tables[7], 6, ["UC-06", "Nộp cả bài", "Người học", "FR-08, FR-09", "Một batch attempt idempotent được lưu."])
    set_row(tables[7], 10, ["UC-10", "Cấu hình StudyLens", "Người học", "FR-14", "Thay đổi loại câu hỏi và độ khó."])
    set_row(tables[8], 5, ["BR-05", "Tắt Extension", "Khi OFF, acquisition/polling dừng, session chưa hoàn tất được đóng và late response bị bỏ qua."])
    set_row(tables[8], 6, ["BR-06", "Transcript fallback", "YouTube caption là nguồn ưu tiên; Gemini URL fallback chỉ áp dụng cho video công khai khi caption thiếu."])
    set_row(tables[8], 7, ["BR-07", "Nguồn câu hỏi", "Câu hỏi phải bám full transcript đã xác thực và có evidence timestamp."])
    set_row(tables[8], 8, ["BR-08", "Số câu", "Mục tiêu là 3/5/10/15 câu tại các mốc <10, <30, <60 và >=60 phút; có thể ít hơn nếu evidence không đủ."])
    set_row(tables[9], 1, ["Extension", "Chrome/Edge Manifest V3, TypeScript, React/Vite, YouTube Timedtext", "Side Panel, persistent activation, full caption acquisition, learning-package polling và seek."])
    set_row(tables[9], 3, ["AI Service", "Python 3.12, FastAPI, google-genai", "Gemini URL transcript fallback, full-video question generation, short-answer grading và schema validation."])
    set_row(tables[9], 4, ["LLM Service", "Google AI Gemini API qua HTTPS", "Phân tích video công khai khi fallback và suy luận ngôn ngữ; chỉ FastAPI gọi."])
    set_row(tables[9], 5, ["Database", "PostgreSQL 16", "Persistence duy nhất cho release, migration EF Core, recovery, backup và restore."])
    set_row(tables[10], 3, ["Video không có caption đủ", "Backend xếp transcriptGenerate; FastAPI yêu cầu Gemini phân tích URL video công khai."])
    set_row(tables[10], 4, ["Gemini transcript không khả dụng", "Giữ error state retryable/non-retryable phù hợp; không công bố quiz hoặc đánh dấu completed."])
    set_row(tables[12], 3, ["Question Count", "Theo duration", "3 / 5 / 10 / 15 là số câu mục tiêu"])
    set_row(tables[12], 6, ["Database", "PostgreSQL", "PostgreSQL 16 qua Docker release"])
    set_row(tables[13], 1, ["Video có caption thủ công", "Timedtext được ưu tiên, full cue giữ timestamp và được Backend xác thực."])
    set_row(tables[13], 3, ["Direct fetch chưa sẵn sàng", "Extension retry metadata trong giới hạn; DOM chỉ là preview trừ khi chứng minh bao phủ toàn video."])
    set_row(tables[13], 4, ["Không có hoặc không đủ caption", "Giữ ON, gửi unavailable/insufficient một lần và kích hoạt Gemini URL fallback cho video công khai."])
    set_row(tables[13], 6, ["Quiz", "Backend gửi full transcript bất biến sang FastAPI; câu hỏi phải có evidence cue/timestamp và không lộ answer key."])

    if not any("Revision 1.2.2" in paragraph.text for paragraph in document.paragraphs):
        document.add_page_break()
        document.add_heading("Revision 1.2.2 Full transcript and full video quiz", level=1)
        document.add_paragraph("Revision này chốt yêu cầu nghiệp vụ cho contract 0.5.0 và thay thế các mô tả timer, segment, quiz 5/10/15 phút hoặc SQLite là database release.")
        document.add_heading("Quy tắc hoàn tất và rollback", level=2)
        add_bullet(document, "Session chỉ completed sau khi một full quiz attempt được chấm và lưu atomically.")
        add_bullet(document, "Lỗi giữa transcript, quiz hoặc grading giữ bản ghi lỗi tối thiểu cho retry nhưng không công bố dữ liệu bán phần.")
        add_bullet(document, "Closing an unfinished context removes provisional capture/quiz data theo contract; legacy data không bị xóa bởi migration.")
        document.add_heading("Số câu mục tiêu", level=2)
        add_revision_table(document, ["Thời lượng video", "Số câu mục tiêu"], [["< 10 phút", "3"], ["10 - < 30 phút", "5"], ["30 - < 60 phút", "10"], [">= 60 phút", "15"]])

    style_tables(document)
    document.save(SRS_PATH)


def update_sds() -> None:
    document = Document(SDS_PATH)
    paragraphs = document.paragraphs

    set_text(paragraphs[4], "Phiên bản 1.2.2 theo contract 0.5.0")
    set_text(paragraphs[5], "02/10/2026")
    set_text(paragraphs[13], "StudyLens AI hỗ trợ học tương tác trên YouTube Web. Khi ON và video hợp lệ, Backend tạo session ngay; Extension lấy full transcript; Backend/FastAPI tạo một quiz cho toàn video và lưu kết quả theo video trong PostgreSQL.")
    set_text(paragraphs[14], "Extension ưu tiên captionTracks và Timedtext JSON3/XML. Backend xác thực hash và cue. Khi caption thiếu, FastAPI dùng Gemini phân tích URL video công khai. Backend xếp durable job cho transcript/quiz, công bố quiz không có answer key và chỉ hoàn tất session sau batch attempt.")
    set_text(paragraphs[16], "Phạm vi dưới đây áp dụng cho MVP contract 0.5.0 trên YouTube Web với PostgreSQL là persistence release.")
    set_text(paragraphs[19], "Lấy full transcript từ captionTracks → Timedtext JSON3 → XML; DOM chỉ là preview hoặc fallback khi chứng minh đủ bao phủ toàn video.")
    set_text(paragraphs[20], "Backend tạo Study Session ngay khi ON và video hợp lệ; sau đó xác thực cue/hash/idempotency, lưu transcript và trạng thái xử lý.")
    set_text(paragraphs[21], "Tạo một quiz cho toàn video với số câu mục tiêu 3/5/10/15 theo thời lượng; không còn timer, interval hoặc segment định kỳ.")
    set_text(paragraphs[22], "FastAPI AI Service gọi Google AI Gemini API để sinh câu hỏi, chấm short answer, tạo giải thích và tạo transcript timestamp từ URL video công khai khi cần fallback; API key chỉ ở server.")
    set_text(paragraphs[23], "Side Panel hiển thị processing state, transcript preview/ready, toàn bộ quiz, batch result, evidence seek và history theo video.")
    set_text(paragraphs[26], "Thu audio/microphone/tabCapture/MediaRecorder hoặc gửi audio từ Extension. Transcript fallback chỉ dùng URL video công khai qua FastAPI/Gemini.")
    set_text(paragraphs[39], "Tách biệt Extension, Backend, FastAPI, Gemini và PostgreSQL. Backend là nguồn sự thật; contract 0.5.0 là seam tích hợp; durable job, idempotency và transaction ngăn dữ liệu trùng hoặc bán phần.")
    set_text(paragraphs[44], "Extension chỉ giao tiếp với ASP.NET Core Backend. Backend tạo session, nhận full transcript, xếp ProcessingJob và polling state từ PostgreSQL. FastAPI là thành phần duy nhất gọi Gemini cho transcript URL fallback, question generation và short-answer grading. PostgreSQL là persistence release duy nhất.")
    set_text(paragraphs[53], "apps/extension/")
    set_text(paragraphs[54], "├── manifest.json, vite.config.ts, scripts/package-extension.ps1")
    set_text(paragraphs[55], "└── src/")
    set_text(paragraphs[56], "    ├── features/video-activation (persistent gate, context, full transcript)")
    set_text(paragraphs[57], "    ├── features/session-quiz (learning package, quiz, retry)")
    set_text(paragraphs[58], "    ├── features/assessment-history (batch attempt, result, history)")
    set_text(paragraphs[59], "    ├── platform/youtube (player, caption tracks, Timedtext, DOM preview)")
    set_text(paragraphs[60], "    ├── shared (HTTP, messaging, validation)")
    set_text(paragraphs[61], "    └── shell (Side Panel composition)")
    set_text(paragraphs[62], "Build output được đóng gói deterministic thành StudyLens-AI-extension-0.5.0.zip cùng SHA256SUMS.txt và release manifest.")
    set_text(paragraphs[64], "Module chạy khi ON và trang /watch hợp lệ. Khi phát hiện video, Extension yêu cầu Backend tạo session ngay, sau đó lấy full Timedtext mà không phụ thuộc currentTime. Abort/stale guard loại response cũ khi A→B, OFF hoặc rời /watch.")
    set_text(paragraphs[67], "4.4 Full Transcript Acquisition")
    set_text(paragraphs[68], "Nguồn ưu tiên là captionTracks → Timedtext JSON3 → XML. Metadata xuất hiện muộn được retry có giới hạn. DOM chỉ được upload khi chứng minh bao phủ gần toàn video; nếu không, gửi insufficient một lần để kích hoạt Gemini fallback.")
    set_text(paragraphs[70], "Manager lưu extensionEnabled trong chrome.storage.local, mặc định OFF lần đầu. Khi ON, mỗi video hợp lệ có context/session riêng; chuyển video giữ ON, đóng context cũ và hủy acquisition/polling stale. OFF dừng mọi tab mà không điều khiển playback.")
    set_text(paragraphs[75], "ASP.NET Core 8 Web API, C#, Entity Framework Core, PostgreSQL 16 và health endpoints.")
    set_text(paragraphs[77], "services/api/src/StudyLens.Api/")
    set_text(paragraphs[78], "├── Features/VideoActivation")
    set_text(paragraphs[79], "├── Features/SessionQuiz (session, transcript, durable worker, retry)")
    set_text(paragraphs[80], "├── Features/AssessmentHistory (batch grading, history)")
    set_text(paragraphs[81], "├── Infrastructure/Persistence (EF Core PostgreSQL, migrations)")
    set_text(paragraphs[82], "├── BuildingBlocks (errors, health, serialization)")
    set_text(paragraphs[83], "└── Program.cs (--migrate-only, API host)")
    set_text(paragraphs[84], "Release container chạy non-root; migration là one-shot job bắt buộc trước API readiness.")
    set_text(paragraphs[96], "services/ai/app/")
    set_text(paragraphs[97], "├── features/transcript_generation")
    set_text(paragraphs[98], "├── features/question_generation")
    set_text(paragraphs[99], "├── features/grading")
    set_text(paragraphs[100], "├── platform/llm (Gemini provider, structured output)")
    set_text(paragraphs[101], "├── platform/config.py")
    set_text(paragraphs[102], "└── main.py and routers")
    set_text(paragraphs[103], "Container chạy non-root, không publish port host và không cài wildcard CORS.")
    set_text(paragraphs[105], "AI Service nhận full transcript đã xác thực để sinh quiz. Riêng fallback, service nhận videoId, tự tạo URL https://www.youtube.com/watch?v={videoId}, yêu cầu Gemini trả full transcript timestamp và kiểm tra Pydantic schema.")
    set_text(paragraphs[107], "AI Service nhận full transcript, questionCount mục tiêu, questionType và difficulty. Transcript dài được chia nội bộ theo giới hạn prompt, sau đó hợp nhất, loại trùng và chọn câu phân bố theo timeline. Kết quả ít hơn mục tiêu được chấp nhận nếu evidence không đủ; quiz rỗng bị từ chối.")
    set_text(paragraphs[109], "MCQ được chấm xác định tại Backend. Short answer được FastAPI/Gemini chấm theo reference answer và evidence. Toàn bộ bài được lưu trong một transaction; provider failure không tạo attempt/history bán phần.")
    set_text(paragraphs[111], "Google AI Gemini API chỉ được FastAPI AI Service gọi qua HTTPS. Backend và Extension không phụ thuộc SDK/model name hay nhận API key.")
    set_text(paragraphs[112], "Provider áp dụng timeout riêng cho transcript, quiz và grading; chuẩn hóa quota/rate limit, timeout, service unavailable và invalid output thành error contract an toàn.")
    set_text(paragraphs[113], "LLM_PROVIDER, GEMINI_API_KEY và GEMINI_MODEL chỉ được inject bằng environment secret khi runtime; không copy vào image, Extension bundle, log hoặc repository.")
    set_text(paragraphs[116], "Entity chính: Videos, StudySessions, TranscriptCaptures, TranscriptCues, ProcessingJobs, Quizzes, QuizQuestions, QuestionOptions, QuizAttempts và AttemptAnswers. Quan hệ giữ identity video xuyên transcript, quiz, attempt và history. TranscriptAudioChunks/AnswerAttempts cũ được giữ legacy khi migration.")
    flow = [
        "1. Người dùng mở video YouTube khi StudyLens ON.",
        "2. Extension xác nhận /watch, videoId và active tab.",
        "3. Backend tạo Study Session active ngay, idempotent theo activation/video.",
        "4. Extension lấy full caption track mà không chờ playback.",
        "5. Extension chuẩn hóa, hash và PUT full transcript một lần.",
        "6. Backend xác thực/lưu; khi insufficient/unavailable thì xếp transcriptGenerate.",
        "7. FastAPI có thể dùng Gemini URL fallback và trả transcript timestamp.",
        "8. Backend xếp quizGenerate cho full transcript ready.",
        "9. FastAPI sinh và validate quiz; Backend publish quiz không answer key.",
        "10. Side Panel polling learning package và hiển thị toàn bộ quiz.",
        "11. Người học nộp full attempt; Backend/FastAPI chấm atomically.",
        "12. PostgreSQL lưu completed session, attempt, answers và history theo video.",
        "13. Side Panel hiển thị result/explanation và seek evidence timestamp.",
        "14. Retry chỉ xếp lại bước failed và không tạo dữ liệu trùng.",
    ]
    for index, text in enumerate(flow, start=121):
        set_text(paragraphs[index], text)
    set_text(paragraphs[138], "ON → persist gate → detect valid video → POST session → acquire full caption → PUT transcript → Backend validation/job → quiz ready → batch attempt → completed/history.")
    set_text(paragraphs[140], "A→B khi ON: abort A, close incomplete A, create B, ignore stale response. Rời /watch: close current context, giữ ON. OFF: broadcast stop, cancel acquisition/polling, close incomplete sessions, không điều khiển playback.")
    set_text(paragraphs[143], "Danh sách endpoint tuân theo public/AI contract 0.5.0 trong contracts/. Health endpoint là operational API và không thay đổi business contract.")
    set_text(paragraphs[149], "Sau batch submit, Side Panel hiển thị score, outcome, reference answer, explanation và Review at Timestamp. Trước submit, response quiz public không chứa correctOptionId hoặc referenceAnswer.")
    set_text(paragraphs[153], "GEMINI_API_KEY, database password và secret không được lưu trong Extension, image, log hay Git repository.")
    set_text(paragraphs[154], "Extension chỉ gọi Backend trên origin được cấu hình; FastAPI không bật wildcard CORS và chỉ nằm trong Docker network.")
    set_text(paragraphs[155], "Docker release chỉ publish Backend trên 127.0.0.1:5000; PostgreSQL và FastAPI không publish port host. Container Backend/AI chạy bằng user không đặc quyền.")
    set_text(paragraphs[156], "Input/output AI, videoId, timestamp, hash và idempotency key được validate trước persistence/publication.")
    set_text(paragraphs[157], "Lỗi server/provider được chuẩn hóa thành envelope an toàn và không làm lộ payload provider, key hay connection string.")
    set_text(paragraphs[165], "Bản release local demo dùng Docker Compose gồm postgres, migrate, ai và api. Job migrate phải hoàn tất trước API; health/live và health/ready phân biệt process liveness với database readiness.")
    set_text(paragraphs[166], "Database release: PostgreSQL 16 với named volume, EF Core migration, backup/restore và restart recovery. SQLite chỉ còn phục vụ một số test legacy, không phải runtime release.")
    set_text(paragraphs[175], "Revision 1.2.2 thiết kế StudyLens AI theo contract 0.5.0: persistent activation, session ngay, full transcript, Gemini public-video fallback, durable jobs, full-video quiz, batch grading, video-linked history và PostgreSQL failure-safe persistence. Docker release cô lập secret/port; Extension được đóng gói deterministic cho Chrome/Edge.")

    tables = document.tables
    set_row(tables[0], 0, ["Tên tài liệu", "Tài liệu Đặc tả Thiết kế Giải pháp StudyLens AI contract 0.5.0"])
    set_row(tables[0], 1, ["Phiên bản", "1.2.2"])
    set_row(tables[0], 2, ["Trạng thái", "Cập nhật kỹ thuật release Week 6"])
    set_row(tables[0], 3, ["Ngày cập nhật", REVISION_DATE])
    add_history_row(tables[1], ["1.2.2", REVISION_DATE, "Nhóm 10", "Contract 0.5.0; full transcript/full-video quiz; Gemini URL fallback; PostgreSQL, durable jobs, Docker deployment, security và recovery."])
    set_table_widths(tables[1], [0.85, 1.4, 1.2, 3.0])
    set_row(tables[2], 6, ["Database", "PostgreSQL 16, Entity Framework Core migrations"])
    set_row(tables[3], 1, ["Extension", "Lưu ON/OFF, cô lập active tab/video, lấy full caption, hiển thị learning package, quiz, result/history và seek."])
    set_row(tables[3], 2, ["Backend", "Nguồn sự thật cho session/transcript/job/quiz/attempt/history; idempotency, transaction, retry và recovery."])
    set_row(tables[3], 3, ["AI Service", "Gemini URL transcript fallback, full-video question generation, short-answer grading và structured-output validation."])
    set_row(tables[3], 5, ["Database", "PostgreSQL lưu business state bền vững; legacy tables được giữ qua migration."])
    data_rows = [
        ["Videos", "VideoId, YoutubeVideoId, Title, DurationMs", "Danh tính video dùng chung cho transcript, quiz, session và history."],
        ["StudySessions", "SessionId, VideoId, Preferences, TranscriptStatus, QuizStatus, Status", "Aggregate trạng thái; chỉ completed sau full attempt."],
        ["TranscriptCaptures/Cues", "Source, Language, ContentHash, Status, StartMs, EndMs, Text", "Full transcript đã xác thực và evidence timestamp."],
        ["ProcessingJobs", "JobType, IdempotencyKey, Status, AttemptCount, LastError", "Durable transcriptGenerate/quizGenerate, tối đa ba lần thử và recovery sau restart."],
        ["Quizzes/Questions/Options", "TranscriptId, ModelName, PromptVersion, Evidence, AnswerKey", "Một quiz toàn video; dữ liệu private không công khai trước submit."],
        ["QuizAttempts/AttemptAnswers", "SessionId, Score, SubmittedAt, Answer, Feedback", "Batch attempt atomically và history theo video/session."],
        ["Legacy", "TranscriptAudioChunks, AnswerAttempts", "Giữ tương thích dữ liệu cũ; không dùng cho luồng 0.5.0 mới."],
    ]
    for row_index, values in enumerate(data_rows, start=1):
        set_row(tables[4], row_index, values)
    set_table_widths(tables[4], [1.7, 2.25, 2.5])
    set_table_font_size(tables[4], 10)
    api_rows = [
        ["POST", "/api/sessions", "Tạo/replay session active ngay cho video hợp lệ."],
        ["PUT", "/api/sessions/{sessionId}/transcript", "Nộp full transcript YouTube hoặc trạng thái unavailable/insufficient."],
        ["GET", "/api/sessions/{sessionId}/learning-package", "Polling transcript, processing error và quiz public."],
        ["POST", "/api/sessions/{sessionId}/retry", "Xếp lại đúng bước transcriptGenerate/quizGenerate failed."],
        ["POST", "/api/sessions/{sessionId}/complete", "Đóng context mà không tự đánh dấu completed."],
        ["POST", "/api/quizzes/{quizId}/attempts", "Nộp và chấm toàn bộ câu trả lời atomically."],
        ["GET", "/api/history", "Lấy lịch sử submitted attempts."] ,
        ["GET", "/api/history/videos/{videoId}", "Lịch sử gắn đúng video."],
        ["POST", "/api/ai/transcripts/generate", "FastAPI tạo timestamp transcript từ URL video công khai."],
        ["POST", "/api/ai/question-generation/generate", "Sinh quiz toàn video với structured output."],
        ["POST", "/api/ai/grading/grade", "Chấm short answer và trả explanation."],
    ]
    for row_index, values in enumerate(api_rows, start=1):
        set_row(tables[5], row_index, values)
    set_table_widths(tables[5], [1.0, 3.0, 2.45])
    set_table_font_size(tables[5], 10)
    set_row(tables[6], 3, ["Không có/không đủ caption", "Giữ session processing; xếp Gemini URL fallback nếu video công khai."])
    set_row(tables[6], 4, ["Gemini transcript không khả dụng", "Lưu error state để retry hoặc hiển thị unavailable; không công bố quiz."])
    set_row(tables[7], 1, ["Extension", "Persistent ON/OFF, active-tab isolation, full transcript acquisition, learning-package polling, quiz/result/history UI và seek."])
    set_row(tables[7], 2, ["Backend", "Session, transcript validation, durable jobs, quiz publication, batch grading, history, migration và health."])
    set_row(tables[7], 3, ["AI Service", "Gemini transcript fallback, full-video question generation, short-answer grading và output validation."])
    set_row(tables[8], 5, ["FR-05", "Processing state", "LearningPackageService + ProcessingJobs + Side Panel polling"])
    set_row(tables[8], 6, ["FR-06", "Full-video learning package", "TranscriptCapture + quizGenerate durable job"])
    set_table_widths(tables[8], [1.1, 2.05, 3.3])
    set_table_font_size(tables[8], 10)

    if not any("Revision 1.2.2" in paragraph.text for paragraph in document.paragraphs):
        document.add_page_break()
        document.add_heading("19 Revision 1.2.2 Release architecture", level=1)
        document.add_paragraph("Phần này chốt thiết kế contract 0.5.0 và thay thế các mô tả session timer, transcript segment, quiz định kỳ và SQLite runtime release.")
        document.add_heading("19.1 Deployment topology", level=2)
        add_revision_table(document, ["Service", "Exposure", "Responsibility"], [
            ["postgres", "Docker network only", "PostgreSQL 16 named volume and health"],
            ["migrate", "One-shot", "Apply EF Core migrations before API"],
            ["ai", "Docker network only", "FastAPI and Gemini provider"],
            ["api", "127.0.0.1:5000", "Public Extension API and health"],
        ])
        document.add_heading("19.2 Recovery invariants", level=2)
        add_bullet(document, "running ProcessingJob được requeue sau Backend restart và không reset attempt/error history.")
        add_bullet(document, "Idempotency key cùng payload hội tụ về một session/capture/job/attempt; payload xung đột bị từ chối.")
        add_bullet(document, "Quiz/attempt/history chỉ publish sau transaction thành công; provider failure không để lại public partial state.")
        document.add_heading("19.3 Release verification", level=2)
        add_bullet(document, "Automated: Extension test/typecheck/build/package; Backend; AI Service; contract; fake-provider release E2E; migration/restart/backup-restore.")
        add_bullet(document, "Manual: M-01…M-08 trên Chrome và Edge với screenshot, sanitized log và database counts.")
        add_bullet(document, "Live provider: ba Gemini cases cho MCQ 3 câu, short answer 5 câu và public-video transcript fallback.")

    replace_diagrams(document)
    style_tables(document)
    document.save(SDS_PATH)


if __name__ == "__main__":
    update_srs()
    update_sds()
    print("Updated Week 6 SRS and SDS documents.")

from __future__ import annotations

import importlib.util
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "docs" / "reports"
DIAGRAMS = ROOT / "artifacts" / "docx-qa" / "final-architecture"
REVISION = "1.2.3"
REVISION_DATE = "13/10/2026"

SRS_SOURCE = ROOT / "SRS_StudyLens_AI_Persistent_Activation_v1.2 (2).docx"
SDS_SOURCE = ROOT / "SDS - StudyLens AI - Học tập có tương tác với Youtube.docx"
SRS_OUTPUT = REPORTS / "SRS_StudyLens_AI_Final_v1.2.3.docx"
SDS_OUTPUT = REPORTS / "SDS_StudyLens_AI_Final_v1.2.3.docx"

SDS_DIAGRAMS = [
    "01-system-context.png",
    "02-logical-architecture.png",
    "03-runtime-architecture.png",
    "04-extension-components.png",
    "05-backend-components.png",
    "06-ai-components.png",
    "07-learning-flow-sequence.png",
    "08-sidepanel-state.png",
    "09-business-capabilities.png",
]


def load_format_helpers():
    path = ROOT / "scripts" / "format-thesis-reports.py"
    spec = importlib.util.spec_from_file_location("study_report_format", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("Không thể nạp bộ định dạng báo cáo.")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


FORMAT = load_format_helpers()


def set_text(paragraph: Paragraph, text: str, style: str | None = None) -> None:
    FORMAT.set_text(paragraph, text)
    if style:
        paragraph.style = style


def set_cell(cell, text: str) -> None:
    FORMAT.set_cell_text(cell, text)


def remove_paragraph(paragraph: Paragraph) -> None:
    parent = paragraph._p.getparent()
    if parent is not None:
        parent.remove(paragraph._p)


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    return FORMAT.paragraph_after(paragraph, text, style)


def insert_paragraph_before(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addprevious(element)
    created = Paragraph(element, paragraph._parent)
    if style:
        created.style = style
    if text:
        run = created.add_run(text)
        FORMAT.set_run_font(run)
    return created


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise RuntimeError(f"Không tìm thấy đoạn bắt đầu bằng: {prefix}")


def append_history_row(table, values: list[str]) -> None:
    row = table.add_row()
    for cell, value in zip(row.cells, values, strict=True):
        set_cell(cell, value)


def add_figure_after(
    anchor: Paragraph,
    image_path: Path,
    caption: str,
    note: str,
    description: str | None = None,
) -> Paragraph:
    cursor = anchor
    if description:
        cursor = insert_paragraph_after(cursor, description, "Normal")
    image_paragraph = insert_paragraph_after(cursor)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.first_line_indent = Cm(0)
    image_paragraph.paragraph_format.space_before = Pt(6)
    image_paragraph.paragraph_format.space_after = Pt(3)
    image_paragraph.paragraph_format.keep_with_next = True
    run = image_paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(15.2))
    caption_paragraph = insert_paragraph_after(image_paragraph, caption, "DATN Figure Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_paragraph = insert_paragraph_after(
        caption_paragraph,
        "Chú giải chức năng: " + note,
        "DATN Figure Note",
    )
    note_paragraph.paragraph_format.first_line_indent = Cm(0)
    return note_paragraph


def replace_inline_images(document: Document, image_names: list[str]) -> None:
    if len(document.inline_shapes) != len(image_names):
        raise RuntimeError(
            f"Số ảnh SDS thay đổi: mong đợi {len(image_names)}, thực tế {len(document.inline_shapes)}."
        )
    for shape, name in zip(document.inline_shapes, image_names, strict=True):
        path = DIAGRAMS / name
        relationship_id, _ = document.part.get_or_add_image(str(path))
        shape._inline.graphic.graphicData.pic.blipFill.blip.embed = relationship_id
        with Image.open(path) as image:
            ratio = image.height / image.width
        width_cm = 12.0 if name == "09-business-capabilities.png" else 15.2
        height_cm = min(width_cm * ratio, 18.8)
        shape.width = Cm(width_cm)
        shape.height = Cm(height_cm)


def set_figure_note(document: Document, caption_prefix: str, note: str) -> None:
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if not paragraph.text.strip().startswith(caption_prefix):
            continue
        paragraph.style = "DATN Figure Caption"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        target = paragraphs[index + 1] if index + 1 < len(paragraphs) else None
        value = "Chú giải chức năng: " + note
        if target is not None and target.text.strip().startswith("Chú giải chức năng:"):
            set_text(target, value, "DATN Figure Note")
        else:
            target = insert_paragraph_after(paragraph, value, "DATN Figure Note")
        target.paragraph_format.first_line_indent = Cm(0)
        return
    raise RuntimeError(f"Không tìm thấy chú thích hình: {caption_prefix}")


def style_document(document: Document, cover_limit: int) -> None:
    FORMAT.configure_styles(document)
    FORMAT.configure_sections(document)
    FORMAT.format_tables(document)
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else ""
        if index < cover_limit:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.2
            paragraph.paragraph_format.space_after = Pt(6)
            for run in paragraph.runs:
                FORMAT.set_run_font(run, 18 if index <= 1 else 13, bold=index <= 1)
            continue
        if style_name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                FORMAT.set_run_font(run)
            continue
        if text.startswith("Bảng "):
            paragraph.style = "DATN Table Caption"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = True
        elif text.startswith("Hình "):
            paragraph.style = "DATN Figure Caption"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
        elif text.startswith("Chú giải chức năng:"):
            paragraph.style = "DATN Figure Note"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_after = Pt(6)
        elif FORMAT.has_drawing(paragraph):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = True
        elif style_name.startswith("TOC"):
            paragraph.paragraph_format.first_line_indent = Cm(0)
        elif style_name.startswith("List"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(1)
        paragraph.paragraph_format.line_spacing = 1.2
        for run in paragraph.runs:
            FORMAT.set_run_font(run, 13)


def update_srs() -> None:
    document = Document(SRS_SOURCE)
    paragraphs = document.paragraphs
    if len(document.tables) != 15:
        raise RuntimeError(f"SRS phải có 15 bảng, thực tế {len(document.tables)}.")

    set_text(
        paragraphs[3],
        "\nPhiên bản: 1.2.3\nNgày: 13/10/2026\nPhạm vi: Kiến trúc cuối cùng, public contract 0.5.0",
    )
    append_history_row(
        document.tables[0],
        [
            REVISION,
            REVISION_DATE,
            "Nhóm StudyLens AI",
            "Đồng bộ kiến trúc nghiệp vụ cuối cùng, PlantUML, prompt 0.5.1, PostgreSQL và trạng thái xác minh.",
        ],
    )

    functions = document.tables[3]
    set_cell(
        functions.rows[4].cells[2],
        "Extension SHALL lấy toàn bộ caption timestamp ngay khi mở video theo thứ tự captionTracks → Timedtext JSON3 → XML. DOM chỉ được dùng khi chứng minh bao phủ toàn video; nếu không đủ, Backend SHALL kích hoạt Gemini URL fallback cho video công khai.",
    )
    set_cell(
        functions.rows[7].cells[2],
        "Backend SHALL gửi full transcript, tiêu đề video đã chuẩn hóa và preference snapshot tới FastAPI. Tiêu đề chỉ là tín hiệu chủ đề; transcript là nguồn bằng chứng duy nhất. Structured output SHALL đúng schema, bám cue và có timestamp hợp lệ.",
    )

    nfr = document.tables[5]
    set_cell(nfr.rows[10].cells[1], "Khả năng vận hành")
    set_cell(
        nfr.rows[10].cells[2],
        "Hệ thống SHALL tách biệt Extension, Backend, AI Service và PostgreSQL; thông tin bí mật chỉ tồn tại phía server và lỗi của một thành phần không được làm mất dữ liệu học tập đã hoàn tất.",
    )

    rules = document.tables[8]
    set_cell(
        rules.rows[7].cells[2],
        "Câu hỏi phải bám full transcript đã xác thực. Tiêu đề video chỉ định hướng nội dung chính, không được dùng làm bằng chứng hoặc làm theo chỉ dẫn chứa trong tiêu đề/transcript.",
    )

    integration = document.tables[9]
    set_cell(
        integration.rows[3].cells[2],
        "Nhận videoTitle và full transcript; tạo transcript dự phòng, quiz toàn video, chấm short answer và kiểm tra schema/cue. Prompt nội bộ hiện hành là 0.5.1.",
    )
    set_cell(
        integration.rows[5].cells[2],
        "Cơ sở dữ liệu chính của hệ thống; lưu video, session, transcript/cue, processing job, quiz, attempt và history.",
    )

    chapter4 = find_paragraph(document, "4. NGƯỜI DÙNG")
    add_figure_after(
        chapter4,
        DIAGRAMS / "10-srs-use-cases.png",
        "Hình 4.1: Sơ đồ ca sử dụng chính của StudyLens AI",
        "Learner là người học điều khiển ON/OFF, làm quiz, xem kết quả, lịch sử và thử lại. YouTube Web cung cấp ngữ cảnh/caption; Google Gemini API chỉ tham gia transcript dự phòng, sinh câu hỏi và chấm câu trả lời ngắn thông qua FastAPI.",
        "Sơ đồ ca sử dụng mô tả các mục tiêu nghiệp vụ mà hệ thống cung cấp cho người học và ranh giới với hai dịch vụ ngoài.",
    )

    data_flow = find_paragraph(document, "11.2. Luồng dữ liệu")
    add_figure_after(
        data_flow,
        DIAGRAMS / "11-srs-main-activity.png",
        "Hình 11.1: Sơ đồ hoạt động của luồng học toàn video",
        "Restore persistent ON/OFF state khôi phục trạng thái bật/tắt; Create or replay study session tạo hoặc dùng lại phiên; Acquire full transcript lấy toàn bộ bản chép lời; Durable quiz job xử lý bền vững; Manual Retry thử lại đúng bước lỗi; Save attempt atomically lưu toàn bộ bài làm trong một transaction.",
        "Sơ đồ hoạt động thể hiện điều kiện kích hoạt, hai nhánh transcript, điều kiện công bố quiz và nguyên tắc không lưu dữ liệu bán phần khi lỗi.",
    )

    chapter14 = find_paragraph(document, "14. YÊU CẦU MỞ RỘNG")
    verification_heading = insert_paragraph_before(chapter14, "13.1. Trạng thái xác minh bản cuối", "Heading 2")
    verification = insert_paragraph_after(
        verification_heading,
        "Kiến trúc và contract trong tài liệu đã được đối chiếu với source code, mô hình dữ liệu và luồng nghiệp vụ hiện tại. Automated tests, typecheck/build và health checks là bằng chứng kỹ thuật riêng; ma trận thủ công M-01…M-08 trên cả Chrome/Edge và ca đánh giá chất lượng Gemini cuối chỉ được ghi nhận đạt khi có evidence thực tế. Tài liệu không suy diễn các hạng mục chưa có bằng chứng.",
        "Normal",
    )
    verification.paragraph_format.first_line_indent = Cm(1)

    srs_replacements = {
        "PostgreSQL là cơ sở dữ liệu bền vững của bản triển khai.": "PostgreSQL là cơ sở dữ liệu bền vững của hệ thống. Backend lưu video, phiên học, transcript, processing job, quiz, bài làm và lịch sử để trạng thái học tập không bị mất khi dịch vụ khởi động lại.",
        "11. PostgreSQL migration": "11. PostgreSQL phải lưu và khôi phục đúng session, transcript, processing job, quiz, attempt và history.",
        "12. API key chỉ ở server": "12. API key chỉ tồn tại phía server và không được đưa vào Extension hoặc dữ liệu trả về người dùng.",
        "Revision này chốt yêu cầu nghiệp vụ cho contract 0.5.0": "Revision này chốt yêu cầu nghiệp vụ cho contract 0.5.0 và thay thế các mô tả cũ về timer, transcript segment và quiz định kỳ.",
        "Closing an unfinished context removes provisional": "Khi một phiên lỗi hoặc bị đóng giữa chừng, hệ thống chỉ giữ trạng thái tối thiểu để thử lại và không công bố transcript, quiz hoặc kết quả bán phần.",
    }
    for paragraph in document.paragraphs:
        for prefix, replacement in srs_replacements.items():
            if paragraph.text.strip().startswith(prefix):
                set_text(paragraph, replacement, paragraph.style.name)
                break

    set_cell(document.tables[12].rows[6].cells[2], "PostgreSQL 16")

    style_document(document, cover_limit=5)
    document.core_properties.title = "StudyLens AI - Software Requirements Specification"
    document.core_properties.subject = "Final architecture, public contract 0.5.0"
    document.core_properties.comments = "Revision 1.2.3 generated from verified repository architecture."
    document.save(SRS_OUTPUT)


def update_sds() -> None:
    document = Document(SDS_SOURCE)
    paragraphs = document.paragraphs
    if len(document.tables) != 9:
        raise RuntimeError(f"SDS phải có 9 bảng, thực tế {len(document.tables)}.")

    set_text(paragraphs[4], "Phiên bản 1.2.3 theo public contract 0.5.0 và prompt nội bộ 0.5.1")
    for index, paragraph in enumerate(paragraphs[:8]):
        if "02/10/2026" in paragraph.text:
            set_text(paragraph, paragraph.text.replace("02/10/2026", REVISION_DATE))
    set_cell(document.tables[0].rows[1].cells[1], REVISION)
    set_cell(document.tables[0].rows[2].cells[1], "Kiến trúc cuối cùng đã đối chiếu source")
    for row in document.tables[0].rows:
        if row.cells[0].text.strip() == "Ngày cập nhật":
            set_cell(row.cells[1], REVISION_DATE)
    append_history_row(
        document.tables[1],
        [
            REVISION,
            REVISION_DATE,
            "Nhóm 10",
            "PlantUML, schema PostgreSQL thực tế, prompt 0.5.1 và trạng thái xác minh cuối.",
        ],
    )

    responsibilities = document.tables[3]
    set_cell(
        responsibilities.rows[3].cells[1],
        "FastAPI stateless nhận videoTitle/full transcript, tạo transcript dự phòng, sinh quiz toàn video, chấm short answer và kiểm tra structured output/cue-index trước khi trả Backend.",
    )
    set_cell(
        responsibilities.rows[5].cells[1],
        "PostgreSQL lưu bền vững video, phiên học, transcript, processing job, quiz, bài làm và lịch sử học tập.",
    )

    database = document.tables[4]
    database_rows = [
        ["Videos", "YoutubeVideoId, Title, DurationMs, CreatedAtUtc, LastSeenAtUtc", "Danh tính video dùng chung cho session, transcript, quiz và history."],
        ["StudySessions", "SessionId, YoutubeVideoId, VideoTitle, Status, TranscriptStatus, QuizStatus", "Aggregate trạng thái; chỉ completed sau khi nộp toàn bộ quiz."],
        ["TranscriptCaptures", "TranscriptCaptureId, SessionId, Source, ContentHash, DurationMs, Status", "Bản transcript toàn video đã được Backend xác thực."],
        ["TranscriptCaptureCues", "TranscriptCueId, TranscriptCaptureId, StartMs, EndMs, Text", "Cue timestamp dùng làm bằng chứng tin cậy cho câu hỏi."],
        ["ProcessingJobs", "ProcessingJobId, JobType, Status, IdempotencyKey, AttemptCount, LastErrorCode", "Durable transcriptGenerate/quizGenerate, retry thủ công và recovery sau restart."],
        ["QuizAssessments", "QuizId, SessionId, YoutubeVideoId, TranscriptCaptureId, ModelName, PromptVersion", "Một quiz toàn video; lưu model và prompt version đã sử dụng."],
        ["QuestionAssessments / QuestionOptions", "QuestionId, Type, Prompt, evidence, answer key, Position", "Câu hỏi, lựa chọn và evidence; đáp án không công khai trước submit."],
        ["QuizAttempts / AttemptAnswers", "ClientAttemptId, QuizId, SessionId, Score, SubmittedAtUtc, Answer, Outcome", "Batch attempt atomically và history theo video/session."],
    ]
    while len(database.rows) < len(database_rows) + 1:
        database.add_row()
    for index, values in enumerate(database_rows, start=1):
        for cell, value in zip(database.rows[index].cells, values, strict=True):
            set_cell(cell, value)

    database_heading = find_paragraph(document, "8.2 Thiết kế bảng")
    legacy_note = insert_paragraph_after(
        database_heading,
        "Hai bảng TranscriptAudioChunks và AnswerAttempts được giữ để tương thích migration/dữ liệu cũ; luồng contract 0.5.0 không dùng chúng làm đường ghi chính.",
        "Normal",
    )
    legacy_note.paragraph_format.first_line_indent = Cm(1)

    api = document.tables[5]
    set_cell(api.rows[11].cells[1], "/api/ai/grading/short-answer")
    set_cell(api.rows[11].cells[2], "FastAPI chấm short answer và trả structured result.")
    for values in [
        ["GET", "/api/health/live", "Liveness của Backend process."],
        ["GET", "/api/health/ready", "Readiness của Backend và PostgreSQL."],
        ["GET", "/api/health/ai", "Tình trạng kết nối Backend tới AI Service."],
        ["GET", "/health", "Health endpoint nội bộ của FastAPI."],
    ]:
        row = api.add_row()
        for cell, value in zip(row.cells, values, strict=True):
            set_cell(cell, value)

    ai_heading = find_paragraph(document, "6.4 Sinh câu hỏi")
    ai_description = insert_paragraph_after(
        ai_heading,
        "Backend gửi videoTitle đã chuẩn hóa tối đa 512 ký tự cùng full transcript. Prompt 0.5.1 xem tiêu đề là tín hiệu xác định chủ đề, còn cue transcript là nguồn sự thật duy nhất; các chunk ngoại vi có thể trả mảng rỗng. Việc bổ sung tiêu đề không làm tăng số lượt gọi Gemini. Kết quả chỉ được chấp nhận khi schema, loại câu hỏi, tính duy nhất và cue-index đều hợp lệ.",
        "Normal",
    )
    ai_description.paragraph_format.first_line_indent = Cm(1)

    ui_heading = find_paragraph(document, "11.1 Side Panel")
    ui_description = insert_paragraph_after(
        ui_heading,
        "Side Panel có ba tab Học tập, Lịch sử và Cài đặt; bố cục responsive cho panel hẹp/rộng. Tab Học tập hiển thị video, tiến độ, transcript đã xác thực, trạng thái xử lý và quiz. Câu hỏi dùng card, vùng chọn cùng hàng với nhãn, màu chủ đạo #008000 và không lộ đáp án trước khi nộp. Tab Lịch sử nhóm attempt theo video/session và cho phép xem lại evidence timestamp.",
        "Normal",
    )
    ui_description.paragraph_format.first_line_indent = Cm(1)

    physical_heading = find_paragraph(document, "3.1.2 Physical Architecture")
    set_text(physical_heading, "3.1.2 Runtime Architecture", "Heading 3")
    next_element = physical_heading._p.getnext()
    while next_element is not None and not next_element.tag.endswith("}p"):
        next_element = next_element.getnext()
    if next_element is None:
        raise RuntimeError("Không tìm thấy mô tả sau mục Runtime Architecture.")
    physical_description = Paragraph(next_element, physical_heading._parent)
    set_text(
        physical_description,
        "Kiến trúc vận hành mô tả các thành phần chính khi hệ thống phục vụ người học. Extension chạy trong Chrome/Edge, đọc ngữ cảnh YouTube và gửi yêu cầu tới Backend. Backend điều phối phiên học, transcript, quiz, chấm điểm và lịch sử; AI Service xử lý nghiệp vụ AI; PostgreSQL lưu dữ liệu bền vững.",
        "Normal",
    )
    set_text(find_paragraph(document, "Hình 3.3:"), "Hình 3.3: Runtime Architecture của StudyLens AI", "DATN Figure Caption")

    backend_replacements = {
        "├── Infrastructure/Persistence": "├── Infrastructure/Persistence (Entity Framework Core và PostgreSQL)",
        "└── Program.cs": "└── Program.cs (application wiring)",
        "Release container chạy non-root": "Backend tập trung các module nghiệp vụ và sử dụng PostgreSQL làm nguồn dữ liệu bền vững.",
        "Container chạy non-root": "AI Service không lưu trạng thái nghiệp vụ; mọi kết quả phải được kiểm tra schema trước khi trả về Backend.",
        "Extension chỉ gọi Backend trên origin": "Extension chỉ gọi Backend qua API đã cấu hình; API key Gemini chỉ tồn tại phía server và không xuất hiện trong Extension.",
        "Docker release chỉ publish Backend": "Backend xác thực dữ liệu đầu vào, AI Service kiểm tra structured output và PostgreSQL lưu trạng thái học tập đã hoàn tất.",
    }
    for paragraph in document.paragraphs:
        for prefix, replacement in backend_replacements.items():
            if paragraph.text.strip().startswith(prefix):
                set_text(paragraph, replacement, paragraph.style.name)
                break

    deployment_heading = find_paragraph(document, "16 Thiết kế triển khai")
    set_text(deployment_heading, "16 Kiến trúc vận hành hệ thống", "Heading 1")
    deployment_note = insert_paragraph_after(
        deployment_heading,
        "Extension là điểm tương tác với người học và chỉ giao tiếp với Backend. Backend cung cấp các nghiệp vụ quản lý video/phiên học, transcript, quiz, chấm điểm và lịch sử; gọi AI Service khi cần xử lý bằng Gemini; đồng thời đọc và ghi dữ liệu nghiệp vụ trong PostgreSQL.",
        "Normal",
    )
    deployment_note.paragraph_format.first_line_indent = Cm(1)

    test_heading = find_paragraph(document, "14 Thiết kế kiểm thử")
    test_status = insert_paragraph_after(
        test_heading,
        "Trạng thái xác minh phải tách biệt giữa thiết kế, automated test và nghiệm thu thật. Source, contract, mô hình dữ liệu và luồng nghiệp vụ đã được đối chiếu trong revision này. Các ca M-01…M-08 trên Chrome/Edge và smoke chất lượng Gemini cuối chỉ được đánh dấu PASS khi có ảnh/log đã làm sạch; nếu chưa có evidence thì giữ trạng thái pending, không suy diễn từ test giả lập.",
        "Normal",
    )
    test_status.paragraph_format.first_line_indent = Cm(1)

    replace_inline_images(document, SDS_DIAGRAMS)
    notes = {
        "Hình 3.1:": "StudyLens Extension thu thập ngữ cảnh/caption và hiển thị Side Panel; ASP.NET Core Backend là nguồn sự thật nghiệp vụ; FastAPI là biên AI stateless duy nhất gọi Gemini; PostgreSQL lưu trạng thái bền vững.",
        "Hình 3.2:": "Presentation Layer chứa Side Panel và YouTube Context Adapter; Extension Application Layer quản lý activation/tab/polling; Backend Application Layer quản lý session/job/assessment; Integration Layer cung cấp REST và FastAPI; Data and External Services gồm PostgreSQL, YouTube Timedtext và Gemini.",
        "Hình 3.3:": "Learner Environment chứa Chrome/Edge và StudyLens Extension. Backend điều phối Session Management, Transcript Processing, Quiz and Grading cùng Learning History; PostgreSQL lưu dữ liệu; AI Service gọi Gemini cho các tác vụ AI.",
        "Hình 4.1:": "Side Panel UI là giao diện; Service Worker điều phối trạng thái và REST; Content Script đọc trang; Video Context Reader đọc ID/title/duration; Transcript Adapter lấy full caption; Player Controller chỉ seek theo thao tác người dùng; chrome.storage.local giữ ON/OFF và preference.",
        "Hình 5.1:": "Minimal API Endpoints nhận request; các application service thực thi nghiệp vụ; Durable Processing Worker claim/retry/recover job; AI HTTP Clients gọi FastAPI; Entity Framework Core thực hiện transaction với PostgreSQL.",
        "Hình 6.1:": "FastAPI Routers phân tuyến tới ba feature service; Prompt Builder 0.5.1 tạo yêu cầu có videoTitle và evidence; Gemini Provider Adapter gọi API; Pydantic Validation chỉ trả DTO hợp lệ hoặc lỗi an toàn.",
        "Hình 9.1:": "Sequence thể hiện một lần nộp full transcript, nhánh Gemini URL fallback khi thiếu caption, durable quiz job, polling learning package, nộp toàn bộ bài và lưu history atomically.",
        "Hình 11.1:": "Off, Acquiring, Processing, Ready, Error, Submitted và History là các trạng thái chính. Manual Retry dùng lại session/job/idempotency key; ba tab Study, History và Settings tổ chức nội dung Side Panel.",
        "Hình 16.1:": "StudyLens Extension gửi thao tác học tập tới Backend API. Các module Video and Session, Transcript Management, Quiz Processing, Assessment và Learning History dùng PostgreSQL; các nghiệp vụ AI đi qua AI Service và Google Gemini API.",
    }
    for caption, note in notes.items():
        set_figure_note(document, caption, note)

    set_text(find_paragraph(document, "Hình 16.1:"), "Hình 16.1: Liên kết nghiệp vụ và dữ liệu của StudyLens AI", "DATN Figure Caption")

    for paragraph in list(document.paragraphs):
        value = paragraph.text.strip()
        if value.startswith("Bản release local demo dùng Docker Compose"):
            remove_paragraph(paragraph)
        elif value.startswith("Database release: PostgreSQL 16"):
            remove_paragraph(paragraph)

    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("Revision 1.2.2 thiết kế StudyLens AI"):
            set_text(
                paragraph,
                "Revision 1.2.3 mô tả StudyLens AI theo contract 0.5.0: persistent activation, session ngay, full transcript, Gemini public-video fallback, durable processing, full-video quiz, batch grading, video-linked history và PostgreSQL bền vững.",
                "Normal",
            )

    history_table = document.tables[1]
    for row in history_table.rows:
        if row.cells[0].text.strip() == "1.2.2":
            set_cell(row.cells[3], "Contract 0.5.0; full transcript/full-video quiz; Gemini URL fallback; PostgreSQL, durable jobs, bảo mật và khôi phục.")
    set_cell(document.tables[2].rows[6].cells[1], "PostgreSQL 16, Entity Framework Core")
    set_cell(document.tables[7].rows[2].cells[1], "Session, transcript validation, processing jobs, quiz publication, batch grading, history và health.")

    style_document(document, cover_limit=6)
    # Bảng schema có nhiều tên trường dài; dùng cỡ 9,5 pt để tránh đẩy đúng
    # một hàng cuối sang trang riêng trong khi vẫn đọc rõ khi in A4.
    for row in database.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    FORMAT.set_run_font(run, 9.5)
    document.core_properties.title = "StudyLens AI - Software Design Specification"
    document.core_properties.subject = "Final architecture, public contract 0.5.0, prompt 0.5.1"
    document.core_properties.comments = "Revision 1.2.3 generated from verified repository architecture."
    document.save(SDS_OUTPUT)


def main() -> None:
    for source in (SRS_SOURCE, SDS_SOURCE):
        if not source.exists():
            raise FileNotFoundError(source)
    for name in SDS_DIAGRAMS + ["10-srs-use-cases.png", "11-srs-main-activity.png"]:
        if not (DIAGRAMS / name).exists():
            raise FileNotFoundError(DIAGRAMS / name)
    REPORTS.mkdir(parents=True, exist_ok=True)
    update_srs()
    update_sds()
    print(SRS_OUTPUT)
    print(SDS_OUTPUT)


if __name__ == "__main__":
    main()
